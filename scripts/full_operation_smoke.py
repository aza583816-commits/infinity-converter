from __future__ import annotations
import io, json, tarfile, zipfile, gzip, bz2, lzma, tempfile
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import pymupdf
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfReader, PdfWriter

from config.settings import settings
from core.storage import TempWorkspace
from core.tooling import TOOLS
from security.file_guard import validate_upload
from converters.engine import ConversionEngine

ROOT=Path(__file__).resolve().parents[1]

class Upload:
    def __init__(self,name,data,mimetype='application/octet-stream'):
        self.filename=name; self.mimetype=mimetype; self._io=io.BytesIO(data)
    def read(self,n=-1): return self._io.read(n)


def make_fixtures(base: Path):
    base.mkdir(parents=True,exist_ok=True)
    # OCR-friendly image
    img=Image.new('RGB',(1400,900),'white'); d=ImageDraw.Draw(img)
    font_path='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    font=ImageFont.truetype(font_path,54) if Path(font_path).exists() else None
    lines=['INFINITY CONVERTER','Invoice 12345','Email test@example.com','URL https://example.com','Total 1250.50','Name Amount','Alice 100','Bob 200']
    y=80
    for line in lines:
        d.text((80,y),line,fill='black',font=font); y+=85
    img.save(base/'image.png'); img.save(base/'image.jpg',quality=95); img.save(base/'image.webp')
    img.save(base/'image.bmp'); img.save(base/'image.tiff')
    # PDF with text/image/link, 3 pages
    doc=pymupdf.open()
    for idx in range(3):
        page=doc.new_page(width=595,height=842)
        page.insert_text((72,90),f'INFINITY CONVERTER Page {idx+1}',fontsize=18)
        page.insert_text((72,130),'Sensitive SECRET Email test@example.com Total 1250.50',fontsize=12)
        if idx==0:
            rect=pymupdf.Rect(72,160,270,190); page.insert_text((72,178),'https://example.com',fontsize=11)
            page.insert_link({'kind':pymupdf.LINK_URI,'from':rect,'uri':'https://example.com'})
            page.insert_image(pymupdf.Rect(72,220,300,366),filename=str(base/'image.png'))
    doc.save(base/'doc.pdf'); doc.close()
    # encrypted PDF
    out=io.BytesIO(); w=PdfWriter(); w.add_blank_page(width=200,height=200); w.encrypt('secret'); w.write(out); (base/'locked.pdf').write_bytes(out.getvalue())
    # office docs
    d=Document(); d.add_heading('Infinity Converter',0); d.add_paragraph('Hello world test@example.com'); tbl=d.add_table(rows=2,cols=2); tbl.cell(0,0).text='Name'; tbl.cell(0,1).text='Value'; tbl.cell(1,0).text='Alice'; tbl.cell(1,1).text='100'; d.save(base/'doc.docx')
    wb=Workbook(); ws=wb.active; ws.append(['name','value']); ws.append(['Alice',100]); ws.append(['Bob',200]); wb.save(base/'sheet.xlsx')
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text='Infinity Converter'; slide.placeholders[1].text='Presentation test'; prs.save(base/'slides.pptx')
    (base/'text.txt').write_text('10\n20\n30\nHello Infinity\ntest@example.com\nhttps://example.com\n',encoding='utf-8')
    (base/'base64.txt').write_text('SGVsbG8gSW5maW5pdHk=',encoding='utf-8')
    (base/'question.txt').write_text('Q: What is 2+2?\nA: 4\n\nQ: Capital of France?\nA: Paris\n',encoding='utf-8')
    (base/'notes.md').write_text('# Infinity\n\nHello **world**\n',encoding='utf-8')
    (base/'data.csv').write_text('name,value\nAlice,100\nBob,200\n',encoding='utf-8')
    (base/'certs.csv').write_text('name\nAlice Example\nBob Example\n',encoding='utf-8')
    (base/'data.json').write_text('[{"name":"Alice","value":100},{"name":"Bob","value":200}]',encoding='utf-8')
    (base/'data.xml').write_text('<root><item name="Alice">100</item></root>',encoding='utf-8')
    (base/'page.html').write_text('<!doctype html><html><head><meta charset="utf-8"></head><body><h1>Infinity</h1><p>Hello world</p></body></html>',encoding='utf-8')
    # zip
    with zipfile.ZipFile(base/'archive.zip','w',zipfile.ZIP_DEFLATED) as z: z.writestr('folder/a.txt','hello'); z.writestr('folder/b.txt','hello'); z.writestr('c.txt','world')
    with tarfile.open(base/'archive.tar','w') as tf: tf.add(base/'text.txt',arcname='text.txt')
    raw=(base/'text.txt').read_bytes()
    with gzip.open(base/'text.gz','wb') as f: f.write(raw)
    (base/'text.bz2').write_bytes(bz2.compress(raw)); (base/'text.xz').write_bytes(lzma.compress(raw))
    with tarfile.open(base/'archive.tar.gz','w:gz') as tf: tf.add(base/'text.txt',arcname='text.txt')
    with tarfile.open(base/'archive.tar.bz2','w:bz2') as tf: tf.add(base/'text.txt',arcname='text.txt')


def choose_fixture(tool, base: Path):
    tid=tool.id
    if tid=='pdf-unlock': return base/'locked.pdf'
    if tid=='bulk-certificate-maker': return base/'certs.csv'
    if tid=='lms-question-bank-formatter': return base/'question.txt'
    if tid=='base64-decode': return base/'base64.txt'
    if tid=='number-list-analyzer': return base/'text.txt'
    if tid in {'tar-gzip-extract','tar-integrity'} and '.gz' in tool.input_ext: return base/'archive.tar.gz'
    if tid=='tar-bzip2-extract': return base/'archive.tar.bz2'
    if '.pdf' in tool.input_ext: return base/'doc.pdf'
    order=['.png','.jpg','.docx','.xlsx','.pptx','.csv','.json','.xml','.html','.md','.txt','.zip','.tar','.gz','.bz2','.xz']
    mapping={'.png':'image.png','.jpg':'image.jpg','.jpeg':'image.jpg','.webp':'image.webp','.bmp':'image.bmp','.tiff':'image.tiff','.docx':'doc.docx','.xlsx':'sheet.xlsx','.pptx':'slides.pptx','.csv':'data.csv','.json':'data.json','.xml':'data.xml','.html':'page.html','.htm':'page.html','.md':'notes.md','.markdown':'notes.md','.txt':'text.txt','.zip':'archive.zip','.tar':'archive.tar','.gz':'text.gz','.tgz':'archive.tar.gz','.bz2':'text.bz2','.tbz2':'archive.tar.bz2','.xz':'text.xz'}
    for ext in order:
        if ext in tool.input_ext: return base/mapping[ext]
    if '.*' in tool.input_ext or '*' in tool.input_ext: return base/'text.txt'
    raise RuntimeError((tid,tool.input_ext))


def options_for(tool):
    opts={f.id:f.default for f in tool.fields}
    replacements={'course':'Engineering','assignment':'Safety Report','student':'Student','quote':'Quality comes from verification.','title':'Completion Certificate','text':'INFINITY','password':'secret'}
    for f in tool.fields:
        if f.required and not opts.get(f.id): opts[f.id]=replacements.get(f.id,'Test')
    return opts

def param_for(tool):
    special={'pdf-extract-pages':'1-2','pdf-delete-pages':'3','pdf-redact':'SECRET','regex-extract':r'\b[A-Za-z]+\b'}
    return special.get(tool.id,tool.param_default or '')


def verify_semantics(tool_id: str, path: Path) -> None:
    """Check security-sensitive outcomes, not only file existence/shape."""
    if tool_id == 'pdf-password-protect':
        reader = PdfReader(str(path), strict=False)
        if not reader.is_encrypted:
            raise RuntimeError('password-protect output is not encrypted')
        if not reader.decrypt('secret'):
            raise RuntimeError('password-protect output does not accept the requested password')
        if len(reader.pages) < 1:
            raise RuntimeError('password-protect output has no readable pages after decryption')
    elif tool_id == 'pdf-unlock':
        reader = PdfReader(str(path), strict=False)
        if reader.is_encrypted:
            raise RuntimeError('unlock output is still encrypted')
        if len(reader.pages) < 1:
            raise RuntimeError('unlock output has no readable pages')


def run():
    engine=ConversionEngine(); results=[]; failures=[]
    with tempfile.TemporaryDirectory(prefix='ic72-smoke-fixtures-') as td:
        base=Path(td); make_fixtures(base)
        for idx,tool in enumerate(TOOLS.values(),1):
            try:
                with TempWorkspace() as ws:
                    safe=[]
                    if tool.input_required:
                        count=2 if tool.id in {'pdf-merge','pdf-compare','text-diff','checksum-compare','zip-create','tar-create','tar-gzip-create','tar-bzip2-create','csv-merge-deduplicate','image-to-pdf'} else 1
                        for n in range(count):
                            fp=choose_fixture(tool,base)
                            data=fp.read_bytes()
                            up=Upload(fp.name,data)
                            item=validate_upload(up,max_bytes=settings.max_file_bytes,inspect_only=False,workspace=ws.path,max_pdf_pages=settings.max_pdf_pages)
                            safe.append(item)
                    res=engine.convert(tool=tool,safe_inputs=safe,workspace=ws,timeout=min(settings.subprocess_timeout,90),max_pdf_pages=settings.max_pdf_pages,param=param_for(tool),options=options_for(tool))
                    if not res.path.exists() or res.path.stat().st_size<=0: raise RuntimeError('empty output')
                    verify_semantics(tool.id, res.path)
                    results.append((tool.id,res.path.suffix,res.mime,res.path.stat().st_size))
                    print(f'OK {idx:03d} {tool.id} -> {res.path.name} {res.path.stat().st_size}')
            except Exception as e:
                failures.append((tool.id,type(e).__name__,str(e)))
                print(f'FAIL {idx:03d} {tool.id}: {type(e).__name__}: {e}')
        print(f'RESULT {len(results)}/{len(TOOLS)} pass')
        if failures:
            print('FAILURES')
            for f in failures: print(f)
            raise SystemExit(1)
if __name__=='__main__': run()
