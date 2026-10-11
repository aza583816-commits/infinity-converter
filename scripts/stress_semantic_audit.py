"""Larger multilingual corpora; preserve all existing semantic assertions.

Native Office renderers can be included explicitly in a runnable environment.
No result from this script certifies those tools or all possible input files.
"""
from __future__ import annotations

import bz2
import csv
import gzip
import hashlib
import json
import lzma
import os
import subprocess
import tarfile
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Pt

from scripts import semantic_output_audit as audit
from scripts.output_input_matrix import diversify

EXCLUDED = frozenset({"word-to-pdf", "excel-to-pdf", "ppt-to-pdf", "txt-to-pdf",
                      "html-to-pdf", "markdown-to-pdf", "csv-to-pdf"})
if os.environ.get("IC_STRESS_INCLUDE_NATIVE") == "1":
    EXCLUDED = frozenset()
ORIGINAL = audit.smoke.make_fixtures
LANGUAGE = os.environ.get("IC_STRESS_LANGUAGE", "mixed")
if LANGUAGE not in {"arabic", "english", "mixed"}:
    raise ValueError(LANGUAGE)


def wording(index: int) -> str:
    english = f"Safety inspection equipment reference {index:05d}"
    arabic = f"فحص السلامة ومعدات الوقاية سجل {index:05d}"
    return {"arabic": arabic, "english": english, "mixed": arabic + " / " + english}[LANGUAGE]


def fixtures(base: Path) -> None:
    ORIGINAL(base)
    diversify(base, "dense")
    # Ten portrait/landscape pages with independently known unique markers.
    with pymupdf.open(base / "doc.pdf") as pdf:
        for index in range(7, 11):
            width, height = (842, 595) if index % 2 else (595, 842)
            page = pdf.new_page(width=width, height=height)
            page.insert_text((40, 45), f"EXTRA PAGE {index}", fontsize=16)
            body = "<div style='font-size:12px'>" + "<br>".join(
                f"{wording(index * 100 + row)} | 00123 | 00999 | 1250.50"
                for row in range(15)) + "</div>"
            spare = page.insert_htmlbox(pymupdf.Rect(40, 70, width - 40, height - 45), body)[0]
            assert spare >= 0, "invalid fixture: PDF body overflow"
        pdf.save(base / "stress.pdf")
    (base / "stress.pdf").replace(base / "doc.pdf")

    doc = Document(base / "doc.docx")
    for index in range(200):
        doc.add_paragraph(wording(index) + " | ID 00123 | amount 1250.50")
    table = doc.add_table(rows=40, cols=4)
    for index, row in enumerate(table.rows):
        for column, cell in enumerate(row.cells):
            cell.text = wording(index * 4 + column) + f" [{index:03d}:{column}]"
    doc.add_paragraph("STRESS FINAL DOCX 00999 / نهاية المستند")
    doc.save(base / "doc.docx")

    workbook = load_workbook(base / "sheet.xlsx")
    for sheet in workbook:
        for index in range(500):
            # Leading zero codes remain strings; numeric amounts remain numbers.
            if sheet.title == "Primary":
                sheet.append([wording(index), f"{index:06d}", index + .125])
            else:
                sheet.append([f"Record {index:05d}", wording(index)])
    workbook.save(base / "sheet.xlsx")
    workbook.close()

    presentation = Presentation(base / "slides.pptx")
    for index in range(len(presentation.slides) + 1, 21):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"STRESS SLIDE {index:02d} / 00123"
        slide.placeholders[1].text = "\n".join(wording(index * 10 + row) for row in range(5))
        # Five bilingual records must fit on the authored slide before export;
        # default 32pt text puts its last line outside the slide canvas.
        for paragraph in slide.placeholders[1].text_frame.paragraphs:
            paragraph.font.size = Pt(24)
    presentation.save(base / "slides.pptx")

    rows = [{"name": wording(i), "code": f"{i:06d}",
             "note": f'quoted "field", comma\nsecond line / {wording(i)}'} for i in range(300)]
    with (base / "data.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    (base / "data.json").write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
    # Retain original known fields required by the existing text oracles.
    with (base / "text.txt").open("a", encoding="utf-8") as stream:
        for i in range(1000):
            stream.write(wording(i) + f" | amount 1250.50 | ID {i:06d}\n")
    with (base / 'notes.md').open('a', encoding='utf-8') as stream:
        for i in range(150):
            stream.write(f'\n## {wording(i)} | 00123 | 00999\n')
    # Add real paragraph elements inside the body, not stray text after the
    # closing HTML element. This keeps every stress record independently
    # meaningful and makes the input valid for different HTML renderers.
    import html
    target = base / 'page.html'
    content = target.read_text(encoding='utf-8')
    paragraphs = ''.join(f'<p>{html.escape(wording(i))} | 00123 | 00999</p>' for i in range(150))
    assert '</body>' in content
    target.write_text(content.replace('</body>', paragraphs + '</body>'), encoding='utf-8')

    # High-resolution image transformations; keep the OCR oracle's English-only
    # recognition fixture separate from multilingual image pixel tests.
    with Image.open(base / "image.png") as source:
        image = source.resize((2400, 2100))
    pen = ImageDraw.Draw(image)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 34)
    pen.rectangle((20, 1800, 2380, 2080), fill="white")
    for i in range(5):
        pen.text((40, 1820 + i * 45), wording(i), fill="black", font=font)
    image.save(base / "stress-image.png")
    image.close()
    # Refresh compressed sources after growing the text, preserving exact
    # source/round-trip relationships rather than accepting stale archives.
    raw = (base / "text.txt").read_bytes()
    (base / "text.gz").write_bytes(gzip.compress(raw))
    (base / "text.bz2").write_bytes(bz2.compress(raw))
    (base / "text.xz").write_bytes(lzma.compress(raw))
    for name, mode in (("archive.tar", "w"), ("archive.tar.gz", "w:gz"),
                       ("archive.tar.bz2", "w:bz2")):
        with tarfile.open(base / name, mode) as archive:
            archive.add(base / "text.txt", arcname="text.txt")


def main() -> None:
    original_choose = audit.smoke.choose_fixture
    def choose(tool, base):
        selected = original_choose(tool, base)
        if selected.name == "image.png" and not (tool.id.startswith("ocr-") or tool.id == "image-ocr"):
            return base / "stress-image.png"
        return selected
    audit.smoke.make_fixtures = fixtures
    audit.smoke.choose_fixture = choose
    audit.TOOLS = {key: value for key, value in audit.TOOLS.items() if key not in EXCLUDED}
    selection = set(filter(None, os.environ.get('IC_STRESS_TOOLS', '').split(',')))
    if selection:
        assert selection <= set(audit.TOOLS), selection - set(audit.TOOLS)
        audit.TOOLS = {key: value for key, value in audit.TOOLS.items() if key in selection}
    os.environ["IC_TEST_FIXTURE_PROFILE"] = "baseline"
    report = Path(os.environ["IC_SEMANTIC_REPORT"])
    try:
        audit.run()
    finally:
        if report.exists():
            data = json.loads(report.read_text())
            # Production images deliberately omit Git metadata and the Git
            # executable. Fingerprint tested source in both local and Docker runs.
            try:
                commit = subprocess.check_output(
                    ["git", "rev-parse", "HEAD"], text=True, stderr=subprocess.DEVNULL).strip()
            except (OSError, subprocess.CalledProcessError):
                commit = None
            source_files = ["converters/office.py", "converters/ocr.py",
                            "scripts/semantic_oracles_office.py",
                            "scripts/semantic_output_audit.py", "scripts/stress_semantic_audit.py"]
            data.update({"stress_language": LANGUAGE,
                         "selected_tools": sorted(selection) if selection else None,
                         "excluded_native_office_tools": sorted(EXCLUDED),
                         "full_registry_verified": False,
                         "source_commit": commit,
                         "tested_source_sha256": {path: hashlib.sha256(Path(path).read_bytes()).hexdigest()
                                                   for path in source_files},
                         "fixture_sizes": {"pdf_pages": 10, "docx_paragraphs_added": 200,
                                           "docx_table_cells_added": 160, "xlsx_rows_added_per_sheet": 500,
                                           "pptx_slides": 20, "csv_json_rows": 300,
                                           "text_rows_added": 1000, "image_pixels": [2400, 2100]}})
            report.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
