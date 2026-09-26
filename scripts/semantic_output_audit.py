"""Independent source-to-output assertions over actual engine results.

Complements (does not replace) smoke checks. Only the explicitly listed
tool IDs have a semantic oracle here; unlisted IDs are reported as uncovered,
not mislabeled as quality-certified. Synthetic, non-sensitive fixtures.
"""
from __future__ import annotations

import csv
import gzip
import hashlib
import io
import json
import os
import sys
import tarfile
import tempfile
import zipfile
from pathlib import Path

import pymupdf
from PIL import Image, ImageChops, ImageDraw, ImageFont, ImageOps, ImageStat
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from config.settings import settings
from core.storage import TempWorkspace
from core.tooling import TOOLS
from security.file_guard import validate_upload
from converters.engine import ConversionEngine
from scripts import full_operation_smoke as smoke
from scripts.output_input_matrix import diversify
from scripts.semantic_oracles_extra import ORACLE_IDS, oracle as extra_oracle
from scripts.semantic_oracles_archive_utility import ORACLE_IDS as ARCHIVE_UTILITY_IDS, oracle as archive_utility_oracle
from scripts.semantic_oracles_pdf_image import ORACLE_IDS as PDF_IMAGE_IDS, oracle as pdf_image_oracle
from scripts.semantic_oracles_office import ORACLE_IDS as OFFICE_IDS, oracle as office_oracle
from scripts.semantic_oracles_ocr import ORACLE_IDS as OCR_IDS, oracle as ocr_oracle

PDF_TEXT = "INFINITY CONVERTER"
IMG_IDS = {"image-to-jpg", "image-to-png", "image-to-webp"}
PDF_KEEP = {"pdf-reorder-pages", "pdf-rotate", "pdf-rotate-selected", "pdf-grayscale", "pdf-page-numbers", "pdf-watermark-text", "pdf-compress", "lms-pdf-size-optimizer", "pdf-crop-margins", "pdf-repair", "pdf-booklet"}
PDF_TO_PAGES = {"pdf-to-jpg", "pdf-to-png"}
CSV_JSON = {"csv-to-json", "json-to-csv", "json-to-xlsx", "csv-to-xlsx", "xlsx-to-csv", "xlsx-to-json"}
TEXT_PRESERVING = {"pdf-to-text", "pdf-to-html", "pdf-to-markdown", "docx-to-text", "docx-to-html", "docx-to-markdown", "markdown-to-html", "markdown-to-text", "pptx-to-text", "pptx-to-markdown"}


def downloaded_payload(path: Path) -> Path:
    if path.suffix.lower() != ".zip":
        return path
    # Single operations can also return a ZIP for expected multi-file results.
    # The individual reports should not silently pass based solely on ZIP size.
    with zipfile.ZipFile(path) as zf:
        assert zf.testzip() is None
        assert all(not p.startswith("/") and ".." not in Path(p).parts for p in zf.namelist())
    return path


def independent_oracle(tool_id: str, source: Path | None, output: Path, fixture: Path) -> str | None:
    """Return a description of a verified source-to-output property or None."""
    output = downloaded_payload(output)
    checked = ocr_oracle(tool_id, source, output, fixture)
    if checked is not None:
        return checked
    checked = office_oracle(tool_id, source, output, fixture)
    if checked is not None:
        return checked
    checked = pdf_image_oracle(tool_id, source, output, fixture)
    if checked is not None:
        return checked
    checked = archive_utility_oracle(tool_id, source, output, fixture)
    if checked is not None:
        return checked
    checked = extra_oracle(tool_id, source, output, fixture)
    if checked is not None:
        return checked
    if tool_id == "pdf-merge":
        distinct_second = fixture / "second-merge.pdf"
        assert distinct_second.is_file()
        with pymupdf.open(source) as first, pymupdf.open(distinct_second) as second, pymupdf.open(output) as merged:
            before_first = [page.get_text("text") for page in first]
            before_second = [page.get_text("text") for page in second]
            actual = [page.get_text("text") for page in merged]
            assert len(before_first) >= 3 and len(before_second) == len(before_first) + 1
            assert "SECOND INPUT UNIQUE" in before_second[-1]
            assert actual == before_first + before_second
            expected_sizes = [(page.rect.width, page.rect.height) for pdf in (first,second) for page in pdf]
            assert [(page.rect.width, page.rect.height) for page in merged] == expected_sizes
        return "two distinct PDF inputs merged without lost/duplicated pages, preserving exact source order, text and geometry"
    if tool_id == "pdf-extract-pages":
        with pymupdf.open(str(source)) as before, pymupdf.open(str(output)) as after:
            assert len(before) >= 3 and len(after) == 2
            for i in range(2):
                assert after[i].get_text("text") == before[i].get_text("text"), (i, "Extracted PDF page content differs from source")
                assert abs(after[i].rect.width - before[i].rect.width) < .02
                assert abs(after[i].rect.height - before[i].rect.height) < .02
        return "only the requested two PDF pages retain exact source text, original sequence and physical dimensions"
    if tool_id == "pdf-delete-pages":
        with pymupdf.open(str(source)) as original, pymupdf.open(str(output)) as pdf:
            expected = [p.get_text("text") for p in original]
            assert len(expected) >= 3
            assert [p.get_text("text") for p in pdf] == expected[:2] + expected[3:]
        return "deleted exactly requested third page while preserving every other source page and text"
    if tool_id == "pdf-to-docx":
        doc = Document(str(output))
        import unicodedata
        extracted = " ".join(unicodedata.normalize("NFKC", p.text) for p in doc.paragraphs)
        extracted += " " + " ".join(unicodedata.normalize("NFKC", c.text) for t in doc.tables for row in t.rows for c in row.cells)
        actual = " ".join(extracted.split())
        with pymupdf.open(source) as before:
            expected_page_markers = []
            for page_no, page in enumerate(before, 1):
                source_text = " ".join(unicodedata.normalize("NFKC", page.get_text("text")).split())
                assert source_text, ("Empty source PDF fixture page", page_no)
                # Independently derive a distinctive marker for *each* page.
                # The mixed corpus adds a landscape page labeled LANDSCAPE PAGE
                # FOUR rather than Page 4: never silently drop that source page.
                marker = (f"Page {page_no}" if f"Page {page_no}" in source_text
                          else f"EXTRA PAGE {page_no}" if f"EXTRA PAGE {page_no}" in source_text
                          else "LANDSCAPE PAGE FOUR" if "LANDSCAPE PAGE FOUR" in source_text
                          else None)
                assert marker, ("PDF fixture page has no independently known marker", page_no, source_text[:300])
                expected_page_markers.append(marker)
                assert marker in actual, ("PDF source page missing from DOCX", page_no, marker, actual[:900])
                # All unique source-page fields, not only their page number,
                # must survive the editable Word conversion.
                for field in ("Columns Alpha 00123", "Beta 00999"):
                    if field in source_text:
                        assert field in actual, ("PDF source field lost in DOCX", page_no, field)
                assert marker in actual, ("PDF source page missing from DOCX", page_no, marker, actual[:900])
            assert abs(doc.sections[0].page_width.inches - before[0].rect.width / 72) < .25
        positions = [actual.find(marker) for marker in expected_page_markers]
        assert all(pos >= 0 for pos in positions) and positions == sorted(positions), (
            "PDF source page markers missing or out of order in DOCX", expected_page_markers, positions)
        assert "test@example.com" in actual and "1250.50" in actual
        return "all source PDF page markers including landscape/mixed pages, unique source fields and original order retained as editable Word text"
    if tool_id == "pdf-grayscale":
        with pymupdf.open(str(source)) as original, pymupdf.open(str(output)) as converted:
            assert len(converted) == len(original)
            for a, b in zip(original, converted):
                assert abs(a.rect.width - b.rect.width) <= 1
                assert abs(a.rect.height - b.rect.height) <= 1
                pix = b.get_pixmap(matrix=pymupdf.Matrix(.4, .4), colorspace=pymupdf.csRGB)
                samples = pix.samples
                assert all(abs(samples[i]-samples[i+1]) <= 1 and abs(samples[i+1]-samples[i+2]) <= 1
                    for i in range(0, len(samples), 3))
        return "all pages retain their geometry and are visually grayscale; warning: renderer rasterizes searchable text"
    if tool_id == "pdf-page-numbers":
        # Source page markers and generated page numbers are different: the
        # former already exists on the original, and cannot prove that this
        # tool added any visible number. Compare actual output word geometry.
        with pymupdf.open(source) as before, pymupdf.open(output) as after:
            assert len(before) == len(after) and len(after) > 1
            for page_no, (old, new) in enumerate(zip(before, after), 1):
                assert old.get_text("text").strip() in new.get_text("text")
                strip_top = (new.rect.y0, new.rect.y0 + 36)
                strip_bottom = (new.rect.y1 - 36, new.rect.y1)
                expected_zone = strip_bottom  # default selected position
                old_words = {(word[4], round(word[0]), round(word[1]))
                             for word in old.get_text("words")}
                numbered = [
                    word for word in new.get_text("words")
                    if word[4] == str(page_no)
                    and expected_zone[0] <= word[1] <= word[3] <= expected_zone[1]
                    and abs((word[0] + word[2]) / 2 - new.rect.width / 2)
                         < new.rect.width * .15
                    and (word[4], round(word[0]), round(word[1])) not in old_words
                ]
                assert len(numbered) == 1, (
                    "Missing, duplicated or misplaced generated PDF page number",
                    page_no, [(w[:5]) for w in new.get_text("words")][-12:],
                )
        return "exactly one NEW visible page number per page in the selected bottom-center area, with original source text intact"
    if tool_id in PDF_KEEP:
        import unicodedata
        norm = lambda value: " ".join(unicodedata.normalize("NFKC", value).split())
        with pymupdf.open(source) as original, pymupdf.open(output) as converted:
            before = [norm(page.get_text("text")) for page in original]
            after = [norm(page.get_text("text")) for page in converted]
            assert before and all(before), "The input fixture must contain text on every page"
            assert after and all(after), "A converted page is blank or has no selectable text"
            if tool_id == "pdf-booklet":
                # A booklet places multiple source pages per sheet and can
                # change the final sheet count and page order.
                joined = " ".join(after)
                # The input corpus may append a landscape page whose source
                # title is LANDSCAPE PAGE FOUR, rather than "Page 4".
                # Derive each page marker from the *actual source page*;
                # never accept a generic marker that was not in that page.
                markers = []
                for number, source_page in enumerate(original, 1):
                    source_text = norm(source_page.get_text("text"))
                    candidates = (f"Page {number}", f"EXTRA PAGE {number}",
                                  "LANDSCAPE PAGE FOUR")
                    marker = next((item for item in candidates
                                   if item in source_text), None)
                    if marker is None:
                        lines = [norm(line) for line in source_page.get_text("text").splitlines()
                                 if norm(line)]
                        assert lines, ("booklet source page has no text marker", number)
                        marker = lines[0]
                    assert marker not in markers, ("booklet markers not unique", number, marker)
                    markers.append(marker)
                    assert marker in joined, ("booklet lost a source page", number, marker)
                assert len(converted) <= len(original) + 1
            else:
                assert len(converted) == len(original), (tool_id, len(original), len(converted))
                for index, (before_text, after_text) in enumerate(zip(before, after), 1):
                    for line in before_text.splitlines():
                        if line.strip():
                            assert norm(line) in after_text, (tool_id, index, line[:300], after_text[:700])
                    # NFKC normalizes text but does not waive a lost line.
                    assert all(norm(line) in after_text for line in original[index - 1].get_text("text").splitlines() if line.strip()), (tool_id, index)
        return "every source PDF page and its original text survives the requested PDF transformation; booklet preserves all source page markers"
    if tool_id == "pdf-unlock":
        locked = PdfReader(str(source))
        assert locked.is_encrypted and locked.decrypt("secret")
        reader = PdfReader(str(output))
        assert not reader.is_encrypted and len(reader.pages) == len(locked.pages)
        for old, new in zip(locked.pages, reader.pages):
            assert abs(float(old.mediabox.width) - float(new.mediabox.width)) < 1
            assert abs(float(old.mediabox.height) - float(new.mediabox.height)) < 1
        return "encrypted source becomes unencrypted and preserves exact page count and geometry"
    if tool_id == "pdf-password-protect":
        before = PdfReader(str(source))
        reader = PdfReader(str(output))
        assert reader.is_encrypted and not reader.decrypt("definitely-not-the-user-password")
        assert reader.decrypt("secret") and len(reader.pages) == len(before.pages)
        for old, new in zip(before.pages, reader.pages):
            assert (new.extract_text() or "").strip() == (old.extract_text() or "").strip()
        return "incorrect password denied; correct password unlocks every original PDF page and exact text"
    if tool_id in IMG_IDS:
        with Image.open(source) as before, Image.open(output) as after:
            expected = ImageOps.exif_transpose(before).convert("RGB")
            actual = after.convert("RGB")
            assert expected.size == actual.size
            assert after.format == {"image-to-jpg":"JPEG", "image-to-png":"PNG", "image-to-webp":"WEBP"}[tool_id]
            # Compare the *actual* foreground where source ink exists; a blank
            # white image would otherwise pass a whole-image mean error on a
            # white-backed document with relatively little black text.
            source_gray = expected.convert("L")
            foreground = Image.new("L", expected.size, 0)
            foreground.paste(255, mask=source_gray.point(lambda v: 255 if v < 160 else 0))
            assert foreground.getbbox(), "the image fixture must contain visible foreground"
            diff = ImageChops.difference(expected, actual)
            ink_error = max(ImageStat.Stat(diff.getchannel(ch), mask=foreground).mean[0] for ch in ("R", "G", "B"))
            assert ink_error < (1 if tool_id == "image-to-png" else 30), (tool_id, ink_error)
        return "requested encoded image format, dimensions and actual dark foreground pixels match source within lossless/lossy tolerance"
    if tool_id == "image-grayscale":
        with Image.open(source) as original, Image.open(output) as img:
            expected = ImageOps.grayscale(ImageOps.exif_transpose(original).convert("RGB"))
            actual = img.convert("RGB")
            assert actual.size == expected.size
            channels = actual.split()
            assert ImageChops.difference(channels[0], channels[1]).getbbox() is None
            assert ImageChops.difference(channels[1], channels[2]).getbbox() is None
            actual_luma = actual.convert("L")
            error = ImageStat.Stat(ImageChops.difference(expected, actual_luma)).mean[0]
            assert error < 5, ("grayscale output lost source visual content", error)
            assert ImageStat.Stat(actual_luma).stddev[0] > 1
        return "every output pixel is grayscale and source visual information survives independent luminance comparison"
    if tool_id == "file-hash":
        expected=hashlib.sha256(source.read_bytes()).hexdigest()
        assert expected in output.read_text(encoding="utf-8").lower()
        return "reported SHA-256 matches independent hash of original bytes"
    if tool_id in {"text-to-base64","base64-decode","hex-encode"}:
        import base64
        actual=output.read_text(encoding="utf-8").strip()
        if tool_id=="text-to-base64":
            assert base64.b64decode(actual).decode("utf-8") == source.read_text(encoding="utf-8")
        elif tool_id=="base64-decode":
            assert actual == base64.b64decode(source.read_text(encoding="utf-8").strip()).decode("utf-8")
        else:
            assert bytes.fromhex(actual).decode("utf-8") == source.read_text(encoding="utf-8")
        return "encoded or decoded content is independently byte-exact"
    if tool_id=="csv-to-json":
        with source.open(encoding="utf-8",newline="") as f: expected=list(csv.DictReader(f))
        actual=json.loads(output.read_text(encoding="utf-8"))
        assert actual==expected
        return "every CSV cell, row and header preserved as JSON"
    if tool_id=="json-to-csv":
        expected=json.loads(source.read_text(encoding="utf-8"))
        with output.open(encoding="utf-8",newline="") as f: actual=list(csv.DictReader(f))
        assert len(actual)==len(expected) and all(all(str(v)==actual[i][k] for k,v in row.items()) for i,row in enumerate(expected))
        return "JSON row values retained when serialized into CSV"
    if tool_id=="xlsx-to-json":
        source_book=load_workbook(source,read_only=True,data_only=True)
        actual=json.loads(output.read_text(encoding="utf-8"))
        try:
            assert set(actual)==set(source_book.sheetnames)
            for sheet in source_book:
                rows=list(sheet.values)
                keys=[str(x).strip() for x in rows[0]]
                assert actual[sheet.title] == [dict(zip(keys,row)) for row in rows[1:]]
        finally: source_book.close()
        return "all sheets, records and typed cells retained in JSON"
    if tool_id=="csv-to-xlsx":
        with source.open(encoding="utf-8",newline="") as f: expected=list(csv.reader(f))
        wb=load_workbook(output,read_only=True)
        try:
            actual=[["" if cell is None else str(cell) for cell in row] for row in wb.active.values]
            assert actual==expected
        finally: wb.close()
        return "CSV row/column boundaries and string values retained in XLSX"
    if tool_id=="xlsx-to-csv":
        wb=load_workbook(source,read_only=True,data_only=True)
        try: expected=[["" if v is None else str(v) for v in row] for row in wb.active.values]
        finally: wb.close()
        with output.open(encoding="utf-8-sig",newline="") as f: actual=list(csv.reader(f))
        assert actual==expected
        return "first sheet values retain their row and column boundaries"
    if tool_id=="text-clean":
        source_lines=source.read_text(encoding="utf-8").replace("\r\n","\n").replace("\r","\n").splitlines()
        expected="\n".join(" ".join(line.split()) for line in source_lines).strip()+"\n"
        actual=output.read_text(encoding="utf-8")
        assert actual==expected, (expected, actual)
        return "whitespace normalized without losing or reordering original tokens"
    if tool_id=="docx-to-text":
        word = Document(source)
        expected = []
        for block in word.iter_inner_content():
            if hasattr(block, "rows"):
                for row in block.rows:
                    expected.append("\t".join(cell.text.strip() for cell in row.cells))
            elif block.text.strip():
                expected.append(block.text)
        actual = output.read_text(encoding="utf-8").splitlines()
        assert actual == expected, ("Word-to-text changed paragraph/table values or original reading order", expected, actual)
        return "all Word paragraphs and every table cell retained in original reading order, including trailing paragraphs and bilingual content"
    if tool_id=="pptx-to-text":
        from pptx import Presentation
        slides = Presentation(source)
        text = output.read_text(encoding="utf-8")
        expected_sections = []
        for index, slide in enumerate(slides.slides, 1):
            values = [shape.text.strip() for shape in slide.shapes
                      if hasattr(shape, "text") and shape.text.strip()]
            expected_sections.append(f"--- Slide {index} ---\n" + "\n".join(values))
        assert text == "\n\n".join(expected_sections)
        return "every original presentation slide and text-bearing shape survives exactly in source slide order"
    if tool_id=="word-to-pdf":
        import unicodedata
        word = Document(source)
        expected = []
        for block in word.iter_inner_content():
            if hasattr(block, "rows"):
                expected.extend(cell.text.strip() for row in block.rows for cell in row.cells)
            elif block.text.strip():
                expected.append(block.text.strip())
        with pymupdf.open(output) as pdf:
            assert len(pdf) >= 1
            page_count = len(pdf)
            assert all(page.get_text("words") for page in pdf), "Blank Word PDF page"
        # PyMuPDF can reorder Arabic glyphs in a valid Unicode PDF. Read the
        # *same output* using an independent logical-text parser, rather than
        # dropping Arabic assertions or tolerating changed source strings.
        parsed = PdfReader(str(output))
        assert len(parsed.pages) == page_count, "Independent PDF parsers disagree on page count"
        text = " ".join(unicodedata.normalize("NFKC", page.extract_text() or "")
                        for page in parsed.pages)
        import re

        def preserve_arabic_numeric_word_boundaries(value: str) -> str:
            # Independent PDF extractors can omit whitespace at script changes
            # despite separate, non-overlapping visible words in the PDF.
            # Restore only Arabic<->digit boundaries. Never split one numeric
            # identifier into two tokens or drop an expected source word.
            return re.sub(
                r"(?<=[\u0600-\u06FF])(?=[0-9])|(?<=[0-9])(?=[\u0600-\u06FF])",
                " ", " ".join(unicodedata.normalize("NFKC", value).split())
            )

        normalized = preserve_arabic_numeric_word_boundaries(text)
        for value in expected:
            wanted = preserve_arabic_numeric_word_boundaries(value)
            # Match the complete source field: a truncated numeric ID cannot
            # pass by matching a prefix of a longer extracted identifier.
            assert re.search(r"(?<!\w)" + re.escape(wanted) + r"(?!\w)", normalized), (
                "Word-to-PDF lost actual original content", value, normalized[:1800]
            )
        return "all source Word paragraph and table cell strings across the entire document survive in PDF Unicode text"
    if tool_id=="pdf-to-text":
        import unicodedata
        text = " ".join(unicodedata.normalize("NFKC", output.read_text(encoding="utf-8")).split())
        previous = -1
        with pymupdf.open(source) as pdf:
            assert len(pdf) >= 3
            for index, page in enumerate(pdf, 1):
                expected = " ".join(unicodedata.normalize("NFKC", page.get_text("text")).split())
                assert expected and expected in text, (index, expected[:400], text[:1200])
                position = text.find(expected, previous + 1)
                assert position >= 0 and position > previous, ("PDF page text reordered", index)
                previous = position
        return "every PDF page's independently extracted full Unicode text survives in original page order"
    if tool_id=="zip-create":
        second = fixture / "second-zip.txt"
        assert second.is_file() and second.read_bytes() != source.read_bytes()
        with zipfile.ZipFile(output) as z:
            assert z.testzip() is None
            members = [entry for entry in z.infolist() if not entry.is_dir()]
            assert len(members) == 2 and len({entry.filename for entry in members}) == 2
            content = [z.read(entry.filename) for entry in members]
            assert content == [source.read_bytes(), second.read_bytes()]
        return "two distinct uploaded files each survive byte-exact in original order as separate ZIP members"
    if tool_id=="gzip-compress":
        assert gzip.decompress(output.read_bytes())==source.read_bytes()
        return "compressed original bytes round-trip exactly"
    if tool_id=="gzip-decompress":
        assert output.read_bytes()==gzip.decompress(source.read_bytes())
        return "decompressed bytes match independent standard library"
    if tool_id=="bzip2-compress":
        import bz2
        assert bz2.decompress(output.read_bytes())==source.read_bytes()
        return "bzip2 content round-trips byte-exact"
    if tool_id=="xz-compress":
        import lzma
        assert lzma.decompress(output.read_bytes())==source.read_bytes()
        return "xz content round-trips byte-exact"
    if tool_id=="json-minify":
        assert json.loads(output.read_text(encoding="utf-8"))==json.loads(source.read_text(encoding="utf-8"))
        return "JSON structural meaning unchanged by minification"
    if tool_id=="uuid-list-generator":
        import uuid
        uuids=[v.strip() for v in output.read_text(encoding="utf-8").splitlines() if v.strip()]
        assert len(uuids)==len(set(uuids)) and len(uuids)>0
        for val in uuids: uuid.UUID(val)
        return "generated UUIDs are valid and unique"
    return None


def run():
    profile=os.environ.get("IC_TEST_FIXTURE_PROFILE","baseline")
    with tempfile.TemporaryDirectory(prefix="ic-semantic-audit-") as directory:
        fixture=Path(directory)/"fixtures"
        smoke.make_fixtures(fixture)
        if profile!="baseline": diversify(fixture,profile)
        failures=[]; verified={}; unverified=[]
        engine=ConversionEngine()
        for index,tool in enumerate(TOOLS.values(),1):
            if tool.id not in {
                "pdf-merge","pdf-extract-pages","pdf-delete-pages","pdf-to-docx",
                "pdf-unlock","pdf-password-protect","image-grayscale","file-hash",
                "text-to-base64","base64-decode","hex-encode","csv-to-json",
                "json-to-csv","xlsx-to-json","csv-to-xlsx","xlsx-to-csv",
                "text-clean","docx-to-text","pptx-to-text","word-to-pdf",
                "pdf-to-text","zip-create","gzip-compress","gzip-decompress",
                "bzip2-compress","xz-compress","json-minify","uuid-list-generator",
            } | IMG_IDS | PDF_KEEP | ORACLE_IDS | ARCHIVE_UTILITY_IDS | PDF_IMAGE_IDS | OFFICE_IDS | OCR_IDS:
                unverified.append(tool.id)
                continue
            with TempWorkspace() as workspace:
                try:
                    inputs=[]
                    original=None
                    if tool.input_required:
                        original=smoke.choose_fixture(tool,fixture)
                        if tool.id == "pdf-annotations-report":
                            annotated=fixture/"annotated.pdf"
                            with pymupdf.open(original) as pdf:
                                note=pdf[0].add_text_annot((72, 400), "QUALITY ANNOTATION")
                                note.set_info(title="Quality Inspector")
                                note.update()
                                pdf.save(annotated)
                            original=annotated
                        elif tool.id in {"ocr-receipt-fields", "ocr-text-deduplicate"}:
                            adapted=fixture/("ocr-receipt.png" if tool.id == "ocr-receipt-fields" else "ocr-duplicates.png")
                            if tool.id == "ocr-receipt-fields":
                                with Image.open(original) as sample:
                                    canvas=sample.convert("RGB")
                                face=ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",54)
                                pen = ImageDraw.Draw(canvas)
                                # Keep the synthetic known-answer amount distinct
                                # from the fixture's existing table/text. A clean
                                # background still exercises real OCR while avoiding
                                # accidental overlap that makes the fixture invalid.
                                y = min(900, max(80, canvas.height - 180))
                                box = pen.textbbox((80, y), "SAR 1250.50", font=face)
                                pen.rectangle(
                                    (box[0] - 5, box[1] - 5, box[2] + 5, box[3] + 5),
                                    fill="white",
                                )
                                pen.text((80, y), "SAR 1250.50", fill="black", font=face)
                            else:
                                canvas=Image.new("RGB",(1400,900),"white")
                                face=ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",54)
                                pen=ImageDraw.Draw(canvas)
                                for y,content in [(80,"INFINITY CONVERTER"),(190,"INFINITY CONVERTER"),(300,"Invoice 12345"),(410,"Invoice 12345")]:
                                    pen.text((80,y),content,fill="black",font=face)
                            canvas.save(adapted)
                            canvas.close()
                            original=adapted
                        elif tool.id == "image-auto-orient":
                            oriented=fixture/"oriented.jpg"
                            with Image.open(original) as raw:
                                exif=raw.getexif()
                                exif[274]=6
                                raw.convert("RGB").save(oriented,format="JPEG",quality=95,exif=exif.tobytes())
                            original=oriented
                        count=2 if tool.id in {"pdf-merge","zip-create","image-to-pdf","checksum-compare","tar-create","tar-gzip-create","tar-bzip2-create","text-diff","csv-merge-deduplicate","pdf-compare"} else 1
                        for n in range(count):
                            member = original
                            if n == 1 and tool.id == "pdf-merge":
                                member = fixture / "second-merge.pdf"
                                with pymupdf.open(original) as second:
                                    extra = second.new_page(width=420, height=600)
                                    extra.insert_text((55, 70), "SECOND INPUT UNIQUE Page 4", fontsize=16)
                                    second.save(member)
                            elif n == 1 and tool.id == "zip-create":
                                member = fixture / "second-zip.txt"
                                member.write_text(
                                    "SECOND INPUT UNIQUE — سلامة 00123\n"
                                    "CSV, PDF and image conversion smoke fixture\n",
                                    encoding="utf-8",
                                )
                            up=smoke.Upload(member.name,member.read_bytes())
                            inputs.append(validate_upload(up,max_bytes=settings.max_file_bytes,inspect_only=False,
                                workspace=workspace.path,max_pdf_pages=settings.max_pdf_pages))
                    requested_param = smoke.param_for(tool)
                    if tool.id == "pdf-reorder-pages":
                        with pymupdf.open(original) as source_pdf:
                            total_pages = len(source_pdf)
                        requested_param = ",".join(map(str, [total_pages, *range(1, total_pages)]))
                    elif tool.id == "image-resize":
                        requested_param = "500"
                    elif tool.id == "image-crop":
                        requested_param = "0,0,100,100"
                    result=engine.convert(tool=tool,safe_inputs=inputs,workspace=workspace,
                        timeout=min(settings.subprocess_timeout,90),max_pdf_pages=settings.max_pdf_pages,
                        param=requested_param,options=smoke.options_for(tool))
                    assertion=independent_oracle(tool.id,original,result.path,fixture)
                    if assertion is None: raise AssertionError("No independent quality assertion assigned")
                    verified[tool.id]=assertion
                    print(f"SEMANTIC PASS {tool.id}: {assertion}",flush=True)
                except Exception as error:
                    failures.append({"id":tool.id,"error":f"{type(error).__name__}: {error}"})
                    print(f"SEMANTIC FAIL {tool.id}: {type(error).__name__}: {error}",flush=True)
        report={"profile":profile,"total":len(TOOLS),"semantic_verified":len(verified),
            "semantic_failures":failures,"unverified":unverified,"verifications":verified,
            "quality_complete":len(verified)==len(TOOLS) and not failures}
        report_path=Path(os.environ.get("IC_SEMANTIC_REPORT","/tmp/semantic-quality-report.json"))
        report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
        print(f"SEMANTIC RESULT {len(verified)}/{len(TOOLS)} verified, {len(failures)} failures, {len(unverified)} not yet independently checked",flush=True)
        if failures: raise SystemExit(1)


if __name__=="__main__": run()
