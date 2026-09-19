"""Run every registered conversion on additional genuinely different input corpora.

The existing one-fixture/operation smoke can miss failures that occur with
Unicode, tables, aspect ratios and multi-page documents. This executable
matrix increases *input coverage*, but cannot by itself prove perfect fidelity
for every possible user document.
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw
from docx import Document
from openpyxl import Workbook

import scripts.full_operation_smoke as smoke


ORIGINAL = smoke.make_fixtures


def diversify(base: Path, profile: str) -> None:
    if profile not in ("mixed", "dense"):
        raise ValueError(f"Unknown fixture profile: {profile}")
    pdf_path = base / "doc.pdf"
    with pymupdf.open(str(pdf_path)) as doc:
        if profile == "mixed":
            page = doc.new_page(width=842, height=595)
            page.insert_text((55, 65), "LANDSCAPE PAGE FOUR", fontsize=19)
            page.insert_text((55, 100), "Columns Alpha 00123 | Beta 00999", fontsize=12)
        else:
            for num in range(4, 7):
                page = doc.new_page(width=595, height=842)
                page.insert_text((55, 70), f"EXTRA PAGE {num} text numeric 00123", fontsize=14)
                for line in range(12):
                    page.insert_text((55, 130 + line * 30), f"Row {line:02d} - Value 00999", fontsize=10)
        temporary = base / "variant.pdf"
        doc.save(str(temporary))
    temporary.replace(pdf_path)

    document = Document(str(base / "doc.docx"))
    document.add_paragraph("Arabic and English: السلامة Safety – code 00123")
    table = document.add_table(rows=3 if profile == "mixed" else 10, cols=3)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f"cell-{i}-{j} / قيمة {i * 10 + j}"
    document.add_paragraph("End marker AFTER TABLE 00999")
    document.save(base / "doc.docx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Primary"
    sheet.append(["name", "code", "amount"])
    for index in range(8 if profile == "mixed" else 32):
        sheet.append(["اسم تجريبي " + str(index), f"00{index:03d}", index + .125])
    secondary = workbook.create_sheet("Secondary")
    secondary.append(["English", "Arabic"])
    secondary.append(["Safety", "سلامة"])
    workbook.save(base / "sheet.xlsx")

    (base / "data.csv").write_text(
        'name,code,note\\n'
        '"اسم، عربي",00123,"contains, comma"\\n'
        '"Line break",00999,"line one\\nline two"\\n',
        encoding="utf-8",
    )
    (base / "data.json").write_text(
        '[{"name":"اسم، عربي","code":"00123","note":"A & B"},'
        '{"name":"Line break","code":"00999","note":"line one\\nline two"}]',
        encoding="utf-8",
    )
    (base / "data.xml").write_text(
        '<root><item name="Arabic">سلامة</item><item name="English">Safety &amp; Fire</item></root>',
        encoding="utf-8",
    )
    (base / "page.html").write_text(
        '<!doctype html><html lang="ar"><head><meta charset="utf-8"></head>'
        '<body><h1>Infinity &amp; السلامة</h1><table><tr><td>Safety</td><td>00123</td></tr></table></body></html>',
        encoding="utf-8",
    )
    (base / "notes.md").write_text(
        '# Infinity / السلامة\\n\\n| Name | Code |\\n|---|---|\\n| Arabic | 00123 |\\n',
        encoding="utf-8",
    )
    (base / "text.txt").write_text(
        '10\\n20\\n30\\nHello Infinity\\ntest@example.com\\nhttps://example.com\\n'
        'Arabic السلامة 00123\\nHELLO Infinity\\n',
        encoding="utf-8",
    )

    width, height = ((500, 210) if profile == "mixed" else (211, 501))
    with Image.open(base / "image.png") as image:
        varied = image.resize((width, height))
        varied.save(base / "image.png")
        varied.convert("RGB").save(base / "image.jpg", quality=92)
        varied.save(base / "image.webp")
        varied.save(base / "image.bmp")
        varied.save(base / "image.tiff")


def matrix_make_fixtures(base: Path) -> None:
    ORIGINAL(base)
    diversify(base, os.environ["IC_TEST_FIXTURE_PROFILE"])


if __name__ == "__main__":
    profile = os.environ.get("IC_TEST_FIXTURE_PROFILE")
    if profile not in ("mixed", "dense"):
        raise SystemExit("Set IC_TEST_FIXTURE_PROFILE=mixed or dense")
    smoke.make_fixtures = matrix_make_fixtures
    print(f"QUALITY INPUT MATRIX PROFILE: {profile}", flush=True)
    smoke.run()
