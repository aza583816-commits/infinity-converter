"""PDF -> Word layout regression: do not silently turn a page into one text paragraph."""
from pathlib import Path

import fitz
import pytest
from docx import Document

from converters.mega_tools import pdf_to_docx


def _make_sample(path: Path, *, scanned: bool = False):
    pdf = fitz.open()
    page = pdf.new_page(width=842, height=595)
    if scanned:
        page.draw_rect(fitz.Rect(40, 40, 300, 220), color=(0, 0, 0), width=2)
    else:
        page.insert_text((50, 60), "COURSE SCHEDULE", fontsize=14)
        page.insert_text((55, 130), "Course", fontsize=10)
        page.insert_text((300, 130), "Time", fontsize=10)
        page.insert_text((55, 175), "Physics", fontsize=10)
        page.insert_text((300, 175), "09:00", fontsize=10)
        for y in (100, 145, 190):
            page.draw_line((40, y), (550, y), color=(0, 0, 0), width=1)
        for x in (40, 280, 550):
            page.draw_line((x, 100), (x, 190), color=(0, 0, 0), width=1)
    pdf.save(str(path))
    pdf.close()


def test_pdf_to_docx_keeps_editable_text_and_page_geometry(tmp_path):
    source = tmp_path / "table.pdf"
    output = tmp_path / "table.docx"
    _make_sample(source)
    pdf_to_docx(source, output)
    result = Document(str(output))
    strings = [p.text for p in result.paragraphs]
    for table in result.tables:
        strings.extend(cell.text for row in table.rows for cell in row.cells)
    combined = " ".join(strings)
    assert "COURSE SCHEDULE" in combined
    assert "Physics" in combined
    assert "09:00" in combined
    assert len(result.sections) >= 1
    assert abs(result.sections[0].page_width.inches - 842 / 72) < 0.15
    assert abs(result.sections[0].page_height.inches - 595 / 72) < 0.15


def test_pdf_to_docx_rejects_scanned_pdf_instead_of_returning_empty_word(tmp_path):
    source = tmp_path / "scanned.pdf"
    output = tmp_path / "scanned.docx"
    _make_sample(source, scanned=True)
    with pytest.raises(ValueError, match="OCR"):
        pdf_to_docx(source, output)
    assert not output.exists()
