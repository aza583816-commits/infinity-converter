"""Simple multilingual spreadsheet PDF fallback must retain cells and sheet names."""
from pathlib import Path
import unicodedata

import pymupdf
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from converters import office


def test_simple_xlsx_pdf_fallback_retains_all_original_cells(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "bilingual.xlsx"
    book = Workbook()
    first = book.active
    first.title = "Primary"
    first.append(["name", "code", "amount"])
    first.append(["اسم تجريبي 0", "00000", 0.125])
    first.append(["اسم تجريبي 1", "00001", 1.125])
    second = book.create_sheet("Secondary")
    second.append(["English", "Arabic"])
    second.append(["Safety", "سلامة"])
    book.save(source)
    output_dir = tmp_path / "pdf"
    def simulated_partial_native_pdf(command, *, timeout, env):
        output = output_dir / (Path(command[-1]).stem + ".pdf")
        with pymupdf.open() as document:
            page = document.new_page()
            page.insert_text((50, 60), "name")
            document.save(output)
    monkeypatch.setattr(office, "_run_libreoffice", simulated_partial_native_pdf)
    rendered = office.office_to_pdf(source, output_dir, timeout=30)
    with pymupdf.open(rendered) as pdf:
        raw = " ".join(page.get_text("text") for page in pdf)
        text = unicodedata.normalize("NFKC", raw)
        for expected in ("Primary", "Secondary", "اسم تجريبي 0", "اسم تجريبي 1",
                         "00000", "00001", "0.125", "1.125", "Safety", "سلامة"):
            assert expected in text, (expected, text)


def test_styled_xlsx_does_not_silently_lose_formatting_on_fallback(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "styled.xlsx"
    book = Workbook()
    book.active["A1"] = "سلامة"
    book.active["A1"].font = Font(bold=True)
    book.save(source)
    output_dir = tmp_path / "pdf"
    def incomplete_native_pdf(command, *, timeout, env):
        output = output_dir / (Path(command[-1]).stem + ".pdf")
        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.save(output)
    monkeypatch.setattr(office, "_run_libreoffice", incomplete_native_pdf)
    with pytest.raises(ValueError, match="تنسيق"):
        office.office_to_pdf(source, output_dir, timeout=30)
