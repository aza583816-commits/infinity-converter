"""Whole cell text remains readable across multi-page Word table boundaries."""
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader

from converters.office import office_to_pdf, _docx_pdf_bidi_spacing_copy


def test_long_table_keeps_each_complete_source_cell(tmp_path):
    doc = Document()
    doc.add_paragraph('TABLE START 00123')
    table = doc.add_table(rows=40, cols=4)
    expected = []
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f'Safety inspection equipment reference {i * 4 + j:05d} [{i:03d}:{j}]'
            expected.append(cell.text)
    doc.add_paragraph('TABLE END 00999')
    source = tmp_path / 'source.docx'
    doc.save(source)
    original = source.read_bytes()
    output = office_to_pdf(source, tmp_path / 'output', 60)
    pages = PdfReader(output).pages
    assert len(pages) >= 3
    text = ' '.join(' '.join((p.extract_text() or '').split()) for p in pages)
    for value in expected:
        assert value in text, value
    assert 'TABLE START 00123' in text and 'TABLE END 00999' in text
    assert source.read_bytes() == original


def test_explicit_source_row_splitting_setting_remains_authoritative(tmp_path):
    doc = Document()
    row = doc.add_table(rows=1, cols=1).rows[0]
    row.cells[0].text = 'Explicit split permission'
    setting = OxmlElement('w:cantSplit')
    setting.set(qn('w:val'), '0')
    row._tr.get_or_add_trPr().append(setting)
    source = tmp_path / 'source.docx'
    doc.save(source)
    assert _docx_pdf_bidi_spacing_copy(source, tmp_path) == source
