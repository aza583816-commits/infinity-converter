"""Independent quality oracles for representative real conversion inputs.

Unlike the 162-operation smoke matrix, these tests assert actual values, ordering,
geometry, pixel properties or security semantics of converted outputs. The corpus
is synthetic and contains no customer uploads or personal records.
"""
from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path

import pytest
import pymupdf
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from openpyxl import Workbook

from converters import images, office_advanced, utility_advanced, pdf, pdf_advanced


def _pdf(path: Path, count: int = 3) -> None:
    doc = pymupdf.open()
    for number in range(1, count + 1):
        page = doc.new_page(width=595, height=842)
        page.insert_text((60, 70), f"UNIQUE PAGE {number}", fontsize=16)
        page.insert_text((60, 110), f"Reference {number * 17}", fontsize=12)
    doc.save(str(path))
    doc.close()


def _pdf_texts(path: Path) -> list[str]:
    with pymupdf.open(str(path)) as doc:
        return [page.get_text() for page in doc]


def test_pdf_merge_preserves_page_content_and_sequence(tmp_path):
    a, b, out = (tmp_path / name for name in ("a.pdf", "b.pdf", "merged.pdf"))
    _pdf(a, 2)
    _pdf(b, 1)
    pdf.merge_pdfs([a, b], out)
    text = _pdf_texts(out)
    assert len(text) == 3
    assert ["UNIQUE PAGE 1", "UNIQUE PAGE 2", "UNIQUE PAGE 1"] == [
        next(line for line in page.splitlines() if line.startswith("UNIQUE PAGE")) for page in text
    ]


def test_pdf_extract_preserves_requested_text_and_page_count(tmp_path):
    source, out = tmp_path / "source.pdf", tmp_path / "extract.pdf"
    _pdf(source)
    pdf.extract_pdf_pages(source, out, "1,3")
    text = _pdf_texts(out)
    assert len(text) == 2
    assert "UNIQUE PAGE 1" in text[0] and "UNIQUE PAGE 3" in text[1]
    assert not any("UNIQUE PAGE 2" in part for part in text)


def test_pdf_reorder_preserves_explicit_user_order(tmp_path):
    source, out = tmp_path / "source.pdf", tmp_path / "reordered.pdf"
    _pdf(source)
    pdf_advanced.reorder_pages(source, out, "3,1,2")
    text = _pdf_texts(out)
    assert [int(page.split("UNIQUE PAGE ")[1][0]) for page in text] == [3, 1, 2]


def test_pdf_rotate_keeps_text_and_only_changes_selected_page(tmp_path):
    source, out = tmp_path / "source.pdf", tmp_path / "rotated.pdf"
    _pdf(source)
    pdf_advanced.rotate_selected(source, out, "2", 90)
    with pymupdf.open(str(out)) as doc:
        assert [page.rotation for page in doc] == [0, 90, 0]
        assert "UNIQUE PAGE 2" in doc[1].get_text()


def test_png_to_jpeg_makes_transparency_policy_explicit(tmp_path):
    source, out = tmp_path / "transparent.png", tmp_path / "flattened.jpg"
    image = Image.new("RGBA", (120, 80), (255, 0, 0, 0))
    image.putpixel((60, 40), (0, 0, 0, 255))
    image.save(source)
    images.convert_image(source, out, "JPEG")
    with Image.open(out) as result:
        assert result.format == "JPEG" and result.mode == "RGB"
        assert result.size == (120, 80)
        assert min(result.getpixel((0, 0))) >= 245
        assert max(result.getpixel((60, 40))) <= 70


@pytest.mark.parametrize("format_name,expected", [("PNG", "PNG"), ("WEBP", "WEBP")])
def test_image_format_conversion_preserves_dimensions(tmp_path, format_name, expected):
    source, out = tmp_path / "source.jpg", tmp_path / ("result." + format_name.lower())
    Image.new("RGB", (124, 76), (10, 100, 210)).save(source, format="JPEG")
    images.convert_image(source, out, format_name)
    with Image.open(out) as result:
        assert result.format == expected and result.size == (124, 76)


def test_image_rotation_swaps_geometry(tmp_path):
    source, out = tmp_path / "source.png", tmp_path / "rotated.png"
    Image.new("RGB", (128, 64), (50, 100, 150)).save(source)
    images.rotate_image(source, out, 90)
    with Image.open(out) as result:
        assert result.size == (64, 128)


def test_image_resize_is_bounded_without_upscaling(tmp_path):
    source, out = tmp_path / "source.png", tmp_path / "small.png"
    Image.new("RGB", (400, 200), (50, 100, 150)).save(source)
    images.resize_image(source, out, 100)
    with Image.open(out) as result:
        assert result.size == (100, 50)


def test_csv_json_csv_round_trip_preserves_unicode_and_record_order(tmp_path):
    source, intermediate, out = [tmp_path / n for n in ("source.csv", "middle.json", "roundtrip.csv")]
    rows = [{"name": "اختبار", "code": "00123", "comment": "a,b"}, {"name": "Alice", "code": "00456", "comment": "line one\nline two"}]
    with source.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["name", "code", "comment"])
        writer.writeheader()
        writer.writerows(rows)
    office_advanced.csv_to_json(source, intermediate)
    assert json.loads(intermediate.read_text(encoding="utf-8")) == rows
    office_advanced.json_to_csv(intermediate, out)
    with out.open(encoding="utf-8", newline="") as stream:
        assert list(csv.DictReader(stream)) == rows


def test_xlsx_to_json_preserves_two_sheets_and_unicode_values(tmp_path):
    source, out = tmp_path / "source.xlsx", tmp_path / "result.json"
    book = Workbook()
    first = book.active
    first.title = "Arabic"
    first.append(["name", "value"])
    first.append(["السلامة", 42])
    second = book.create_sheet("English")
    second.append(["name", "value"])
    second.append(["Physics", 17])
    book.save(source)
    office_advanced.xlsx_to_json(source, out)
    result = json.loads(out.read_text(encoding="utf-8"))
    assert result == {"Arabic": [{"name": "السلامة", "value": 42}], "English": [{"name": "Physics", "value": 17}]}


def test_docx_to_text_preserves_paragraph_and_table_cells(tmp_path):
    source, out = tmp_path / "source.docx", tmp_path / "result.txt"
    document = Document()
    document.add_paragraph("مقدمة عربية")
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip((c for row in table.rows for c in row.cells), ("Course", "Time", "Physics", "09:00")):
        cell.text = value
    document.save(source)
    office_advanced.docx_to_text(source, out)
    result = out.read_text(encoding="utf-8")
    assert "مقدمة عربية" in result
    assert "Course\tTime" in result
    assert "Physics\t09:00" in result


def test_docx_to_html_escapes_untrusted_text_and_keeps_table(tmp_path):
    source, out = tmp_path / "source.docx", tmp_path / "result.html"
    document = Document()
    document.add_paragraph("<script>alert(1)</script>")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "اسم"
    table.cell(0, 1).text = "A & B"
    document.save(source)
    office_advanced.docx_to_html(source, out)
    result = out.read_text(encoding="utf-8")
    assert "<script>" not in result and "&lt;script&gt;" in result
    assert "<table>" in result and "اسم" in result and "A &amp; B" in result


def test_text_utilities_output_expected_cleaned_and_deduplicated_content(tmp_path):
    source, cleaned, unique = [tmp_path / n for n in ("input.txt", "clean.txt", "unique.txt")]
    source.write_text("  Hello   world \r\n  Hello   world \r\n  اختبار  \r\n", encoding="utf-8")
    utility_advanced.clean_text(source, cleaned)
    assert cleaned.read_text(encoding="utf-8") == "Hello world\nHello world\nاختبار\n"
    utility_advanced.deduplicate_text(cleaned, unique)
    assert unique.read_text(encoding="utf-8") == "Hello world\nاختبار\n"


def test_number_list_analysis_has_correct_median_and_mean(tmp_path):
    source, out = tmp_path / "numbers.txt", tmp_path / "results.json"
    source.write_text("1, 2, 3, 10", encoding="utf-8")
    utility_advanced.number_list_analysis(source, out)
    result = json.loads(out.read_text(encoding="utf-8"))
    assert result["count"] == 4
    assert result["sum"] == 16
    assert result["median"] == 2.5
    assert result["mean"] == 4


def test_valid_zip_preserves_nested_unicode_filename_and_bytes(tmp_path):
    from converters.archive import create_zip
    data = tmp_path / "original.txt"
    data.write_text("سلامة\n", encoding="utf-8")
    out = tmp_path / "result.zip"
    create_zip([(data, "nested/ملف.txt")], out)
    with zipfile.ZipFile(out) as archive:
        assert archive.namelist() == ["nested/ملف.txt"]
        assert archive.read("nested/ملف.txt") == data.read_bytes()
