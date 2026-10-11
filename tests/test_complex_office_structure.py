"""Synthetic styled Office files: exact records, formula values, images and order."""
from pathlib import Path
from io import BytesIO
import os
import subprocess
import unicodedata

import pymupdf
from PIL import Image, ImageDraw, ImageColor
from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader

from converters.office import office_to_pdf
from scripts.semantic_oracles_office import _field_present, _logical_text_with_geometric_spaces


def normal(value):
    return ' '.join(unicodedata.normalize('NFKC', value).split())


def diagram(path):
    image = Image.new('RGB', (640, 240), 'white')
    pen = ImageDraw.Draw(image)
    for i, color in enumerate(('#c93664', '#18a9ac', '#efbc32', '#536bce')):
        pen.rectangle((20+i*155, 20, 150+i*155, 215), fill=color)
        pen.ellipse((45+i*155, 70, 120+i*155, 155), fill='white')
    image.save(path)
    image.close()


def convert(source, output_dir, tool):
    output_dir.mkdir(exist_ok=True)
    url = os.environ.get('IC_COMPLEX_LIVE_URL')
    if not url:
        return office_to_pdf(source, output_dir, timeout=120)
    output = output_dir / 'actual.pdf'
    status = subprocess.check_output(['curl', '-sS', '--max-time', '240', '-o', str(output),
        '-w', '%{http_code}', '-F', 'tool='+tool, '-F', 'files=@'+str(source), url], text=True)
    assert status == '200', (tool, status, output.read_bytes()[:1500])
    return output


def pdf_evidence(output):
    reader = PdfReader(output)
    text = normal(' '.join(_logical_text_with_geometric_spaces(p) for p in reader.pages))
    with pymupdf.open(output) as pdf:
        assert len(pdf) == len(reader.pages)
        checked_images=set()
        for number, page in enumerate(pdf, 1):
            assert page.get_text().strip(), ('blank page', number)
            for word in page.get_text('words'):
                assert -2 <= word[0] < word[2] <= page.rect.width+2
                assert -2 <= word[1] < word[3] <= page.rect.height+2
            for embedded in page.get_images():
                for rect in page.get_image_rects(embedded[0]):
                    assert page.rect.contains(rect), ('clipped source diagram',number,rect)
                if embedded[0] in checked_images:
                    continue
                checked_images.add(embedded[0])
                with Image.open(BytesIO(pdf.extract_image(embedded[0])['image'])) as original:
                    decoded=original.convert('RGB').resize((640,240))
                    for i,color in enumerate(('#c93664','#18a9ac','#efbc32','#536bce')):
                        for x,y,expected in ((30+i*155,30,ImageColor.getrgb(color)),
                                             (85+i*155,110,(255,255,255)),
                                             (30+i*155,200,ImageColor.getrgb(color))):
                            assert max(abs(a-b) for a,b in zip(decoded.getpixel((x,y)),expected)) < 20
                    decoded.close()
        images = sum(len(p.get_images()) for p in pdf)
        sizes = [(p.rect.width, p.rect.height) for p in pdf]
    return text, images, sizes


def test_complex_word_sections_merged_tables_images_and_every_record(tmp_path):
    photo = tmp_path/'diagram.png'; diagram(photo)
    doc = Document()
    doc.styles['Normal'].font.size = Pt(10)
    doc.sections[0].header.paragraphs[0].text = 'INFINITY ENGINEERING / تقرير السلامة'
    doc.sections[0].footer.paragraphs[0].text = 'SYNTHETIC QA 00123 / بيانات تجريبية'
    expected = []; starts = []
    for block in range(8):
        if block:
            section = doc.add_section(WD_SECTION_START.NEW_PAGE)
            if block % 2:
                section.orientation = WD_ORIENT.LANDSCAPE
                section.page_width, section.page_height = Inches(11.7), Inches(8.3)
            else:
                section.orientation = WD_ORIENT.PORTRAIT
                section.page_width, section.page_height = Inches(8.3), Inches(11.7)
        start = f'SECTION{block:02d} START / السلامة'; starts.append(start)
        doc.add_heading(start, 1)
        doc.add_picture(str(photo), width=Inches(4))
        table = doc.add_table(rows=15, cols=4);table.style='Table Grid'
        table.cell(0,0).merge(table.cell(0,3)).text = f'MERGED{block:02d} / معدات الوقاية'
        expected.append(f'MERGED{block:02d} / معدات الوقاية')
        for i in range(1,15):
            for j in range(4):
                value = f'W{block:02d}{i:02d}{j}X فحص السلامة Safety {i*10+j}.25'
                table.cell(i,j).text = value;expected.append(value)
        end = f'SECTION{block:02d} END 00999';doc.add_paragraph(end);expected.append(end)
    source=tmp_path/'complex.docx';doc.save(source);before=source.read_bytes()
    output=convert(source,tmp_path/'word-output','word-to-pdf')
    text, images, sizes=pdf_evidence(output)
    assert all(_field_present(s,text) for s in expected), [s for s in expected if not _field_present(s,text)][:8]
    assert images >= 8, ('missing embedded diagrams', images)
    assert any(w>h for w,h in sizes) and any(h>w for w,h in sizes)
    positions=[text.find(normal(s)) for s in starts]
    assert all(p>=0 for p in positions) and positions == sorted(positions)
    assert 'INFINITY ENGINEERING' in text and 'SYNTHETIC QA 00123' in text
    assert source.read_bytes() == before


def test_styled_excel_formulas_merged_headers_chart_images_and_all_rows(tmp_path):
    photo=tmp_path/'diagram.png';diagram(photo)
    book=Workbook();expected=[];row_values=[]
    for s in range(3):
        sheet=book.active if not s else book.create_sheet()
        sheet.title=f'Ledger{s+1}'
        title=f'LEDGER{s+1} / السلامة Safety';sheet.merge_cells('A1:E1');sheet['A1']=title
        sheet['A1'].font=Font(size=16,bold=True);expected.append(title)
        sheet.append(['ID','Safety / السلامة','Qty','Price','Total'])
        for i in range(1,81):
            r=i+2; marker=f'L{s+1}R{i:03d}X';label=f'معدات Safety {i:03d}'
            qty=i%5+1;price=i+.25
            sheet.append([marker,label,qty,price,f'=C{r}*D{r}'])
            sheet.cell(r,4).number_format=sheet.cell(r,5).number_format='0.00'
            expected.extend((marker,label));row_values.append((marker,f'{qty*price:.2f}'))
            sheet.row_dimensions[r].height=28
            for cell in sheet[r]:
                cell.alignment=Alignment(wrap_text=True,vertical='center')
                cell.fill=PatternFill('solid',fgColor='EAF5F4' if i%2 else 'FFFFFF')
        for col,width in [('A',18),('B',32),('C',8),('D',12),('E',14)]:sheet.column_dimensions[col].width=width
        image=SheetImage(str(photo));image.width=320;image.height=120;sheet.add_image(image,'A85')
        chart=BarChart();chart.title=f'CHART{s+1} SAFETY TOTALS'
        chart.add_data(Reference(sheet,min_col=5,min_row=2,max_row=12),titles_from_data=True)
        chart.legend=None
        chart.series[0].graphicalProperties.solidFill='007FA3'
        chart.width=16;chart.height=7;sheet.add_chart(chart,'A94')
        sheet.print_area='A1:E110';sheet.print_title_rows='1:2'
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    source=tmp_path/'complex.xlsx';book.save(source);book.close();before=source.read_bytes()
    output=convert(source,tmp_path/'excel-output','excel-to-pdf')
    text,images,sizes=pdf_evidence(output)
    assert all(_field_present(s,text) for s in expected), [s for s in expected if not _field_present(s,text)][:8]
    assert images >= 3, ('missing embedded diagrams',images)
    for s in range(3):assert _field_present(f'CHART{s+1} SAFETY TOTALS',text)
    # Formula values must appear on the actual geometric row of their unique ID.
    with pymupdf.open(output) as pdf:
        # A title alone is insufficient: each chart must contain all ten
        # formula-driven bars, with heights proportional to independently
        # calculated values. Explicit color distinguishes bars from table fills.
        expected_bars=[(i % 5 + 1) * (i + .25) for i in range(1,11)]
        for s in range(3):
            pages=[p for p in pdf if f'CHART{s+1}' in p.get_text()]
            assert len(pages)==1
            bars=[d['rect'] for d in pages[0].get_drawings()
                  if d.get('fill') and all(abs(a-b)<.002 for a,b in
                     zip(d['fill'],(0,127/255,163/255)))]
            bars.sort(key=lambda box:box.x0)
            assert len(bars)==10, ('missing chart data bars',s+1,len(bars))
            tallest=max(box.height for box in bars)
            for box,value in zip(bars,expected_bars):
                assert abs(box.height/tallest-value/max(expected_bars)) < .01
        for marker,amount in row_values:
            matches=[(page,box) for page in pdf for box in page.search_for(marker)]
            assert len(matches)==1, (marker,len(matches))
            page,box=matches[0]
            row=normal(page.get_text(clip=pymupdf.Rect(0,box.y0-3,page.rect.width,box.y1+3)))
            assert _field_present(amount,row), (marker,amount,row)
    assert all(w>h for w,h in sizes)
    positions=[text.find(f'LEDGER{i}') for i in (1,2,3)]
    assert all(p>=0 for p in positions) and positions==sorted(positions)
    assert source.read_bytes()==before
