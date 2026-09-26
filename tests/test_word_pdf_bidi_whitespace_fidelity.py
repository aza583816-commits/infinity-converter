"""Word PDF regression: Arabic-number cell whitespace remains selectable and visible."""
from __future__ import annotations

import hashlib
import shutil
import unicodedata
from pathlib import Path

import pymupdf
from docx import Document
from pypdf import PdfReader

from converters.office import _docx_pdf_bidi_spacing_copy, office_to_pdf


def _source(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Arabic and English: السلامة Safety – code 00123")
    table = doc.add_table(rows=2, cols=2)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            run = cell.paragraphs[0].add_run(f"cell-{i}-{j} / قيمة {i * 10 + j}")
            run.bold = True
    doc.save(path)


def test_word_pdf_does_not_merge_arabic_text_with_ascii_digits(tmp_path: Path) -> None:
    original = tmp_path / "source.docx"
    _source(original)
    original_bytes = original.read_bytes()
    original_checksum = hashlib.sha256(original_bytes).hexdigest()

    output = office_to_pdf(original, tmp_path / "render", timeout=60)

    with PdfReader(str(output)) as pdf:
        actual = " ".join(page.extract_text() or "" for page in pdf.pages)
        logical = " ".join(unicodedata.normalize("NFKC", actual).split())
    for i in range(2):
        for j in range(2):
            expected = f"cell-{i}-{j} / قيمة {i * 10 + j}"
            assert expected in logical, (expected, logical)
    with pymupdf.open(output) as rendered:
        assert len(rendered) >= 1
        assert any(word[4] == "00123"
                   for page in rendered for word in page.get_text("words"))
        assert all(page.get_text("words") for page in rendered)

    assert hashlib.sha256(original.read_bytes()).hexdigest() == original_checksum
    assert not any((tmp_path / "render").glob(".lo-word-input-*")), (
        "Private DOCX render copy must be deleted after conversion"
    )


def test_word_render_copy_preserves_original_styles_and_document(tmp_path: Path) -> None:
    original = tmp_path / "original.docx"
    _source(original)
    before = original.read_bytes()
    render_copy = _docx_pdf_bidi_spacing_copy(original, tmp_path)
    assert render_copy != original
    try:
        doc = Document(render_copy)
        for row in doc.tables[0].rows:
            for cell in row.cells:
                assert cell.paragraphs[0].runs[0].bold is True
        assert "\u00a0" in doc.tables[0].cell(0, 0).text
        assert original.read_bytes() == before
    finally:
        shutil.rmtree(render_copy.parent, ignore_errors=True)


def test_word_without_arabic_digit_boundary_is_not_rewritten(tmp_path: Path) -> None:
    source = tmp_path / "english.docx"
    doc = Document()
    doc.add_paragraph("Safety report 00123 and table 57")
    doc.save(source)
    assert _docx_pdf_bidi_spacing_copy(source, tmp_path) == source
