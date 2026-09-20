"""Real legacy DOC -> PDF regression, not just DOCX extension acceptance.

An independent LibreOffice import generates a genuine binary Word 97 DOC
from an in-memory modern Word fixture. Then exercise the production converter
on that DOC and check the resulting PDF's selectable source content.
"""
from pathlib import Path
import subprocess

import pymupdf
from docx import Document

from converters.office import office_to_pdf


def test_genuine_legacy_doc_is_rendered_with_all_original_text(tmp_path: Path):
    docx = tmp_path / "legacy-source.docx"
    doc = Document()
    doc.add_heading("LEGACY WORD DOCUMENT 00123", 1)
    doc.add_paragraph("Safety inspection report, 2026.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Asset"
    table.cell(0, 1).text = "Finding"
    table.cell(1, 0).text = "Pump 07"
    table.cell(1, 1).text = "No leaks"
    doc.save(docx)

    convert_dir = tmp_path / "doc"
    convert_dir.mkdir()
    args = [
        "libreoffice", "--headless", "--nologo", "--nodefault",
        "--nofirststartwizard",
        f"-env:UserInstallation={(tmp_path / 'lo-import-profile').resolve().as_uri()}",
        "--convert-to", "doc:MS Word 97",
        "--outdir", str(convert_dir), str(docx),
    ]
    completed = subprocess.run(args, capture_output=True, text=True, timeout=45, check=True)
    legacy = convert_dir / "legacy-source.doc"
    assert legacy.exists() and legacy.stat().st_size > 1000, (
        completed.stdout, completed.stderr
    )
    assert legacy.read_bytes()[:8] == bytes.fromhex("d0cf11e0a1b11ae1"), (
        "Expected a real OLE2 legacy DOC, not DOCX renamed with .doc"
    )

    pdf = office_to_pdf(legacy, tmp_path / "output", timeout=60)
    with pymupdf.open(pdf) as output:
        assert len(output) >= 1
        text = " ".join(" ".join(page.get_text("text").split()) for page in output)
    for expected in (
        "LEGACY WORD DOCUMENT 00123", "Safety inspection report",
        "Asset", "Finding", "Pump 07", "No leaks"
    ):
        assert expected in text, ("Legacy DOC conversion lost text", expected, text)
