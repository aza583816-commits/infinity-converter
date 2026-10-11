"""Private chart pagination must preserve source bytes and all workbook objects."""
import zipfile
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Font
from openpyxl.worksheet.pagebreak import Break
from PIL import Image

from converters.office import _xlsx_pdf_tail_chart_copy


def source_workbook(tmp_path, case='default'):
    book=Workbook();sheet=book.active
    sheet.append(['Title','Amount']);sheet.append(['Safety / السلامة',12.50])
    sheet.append(['A',25]);sheet.append(['B','=B2*2'])
    sheet['A1'].font=Font(bold=True)
    sheet.print_title_rows='1:2';sheet.print_area='A1:E50'
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.page_setup.paperSize=sheet.PAPERSIZE_A4;sheet.page_setup.orientation='landscape'
    chart=BarChart();chart.add_data(Reference(sheet,min_col=2,min_row=2,max_row=4))
    chart.width=16;chart.height=7;sheet.add_chart(chart,'A30')
    if case=='explicit':sheet.row_breaks.append(Break(id=25))
    if case=='inside-data':sheet['A35']='LAST ORIGINAL DATA'
    if case=='oversized':chart.height=30
    if case=='crossing-image':
        path=tmp_path/'photo.png';Image.new('RGB',(160,120),'#b84456').save(path)
        sheet.add_image(SheetImage(path),'A26')
    source=tmp_path/'source.xlsx';book.save(source);book.close();return source


def test_tail_chart_copy_changes_only_print_break_xml_and_keeps_formula_and_styles(tmp_path):
    source=source_workbook(tmp_path);before=source.read_bytes()
    copy=_xlsx_pdf_tail_chart_copy(source,tmp_path)
    assert copy!=source and source.read_bytes()==before
    with zipfile.ZipFile(source) as original,zipfile.ZipFile(copy) as updated:
        assert original.namelist()==updated.namelist()
        changed=[n for n in original.namelist() if original.read(n)!=updated.read(n)]
        assert changed==['xl/worksheets/sheet1.xml']
    book=load_workbook(copy)
    try:
        assert [b.id for b in book.active.row_breaks.brk]==[26]
        assert book.active['B4'].value=='=B2*2'
        assert book.active['A1'].font.bold and len(book.active._charts)==1
    finally:book.close()


@pytest.mark.parametrize('case',['explicit','inside-data','oversized','crossing-image'])
def test_authored_breaks_and_unsafe_chart_placements_are_not_overridden(tmp_path,case):
    source=source_workbook(tmp_path,case);before=source.read_bytes()
    assert _xlsx_pdf_tail_chart_copy(source,tmp_path)==source
    assert source.read_bytes()==before
