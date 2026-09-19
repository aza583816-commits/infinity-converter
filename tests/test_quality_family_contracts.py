"""Independent, source-to-output quality assertions for representative tool families.

These tests deliberately check *meaning*, geometry and content, not just file existence.
The full 162-operation smoke matrix separately exercises every registered route/engine.
Synthetic samples only: never commit customer uploads or identifying information.
"""
from __future__ import annotations

import csv
import gzip
import io
import json
import lzma
import tarfile
import zipfile

import fitz
import pytest
from PIL import Image, ImageChops
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from converters import images, office, office_advanced, pdf, mega_tools


def _make_pdf(path, title, width=595, height=842):
    doc = fitz.open()
    page = doc.new_page(width=width, height=height)
    page.insert_text((60, 90), title, fontsize=14)
    doc.save(str(path))
    doc.close()


def _pdf_text(path):
    with fitz.open(str(path)) as doc:
        return [page.get_text("text") for page in doc]


def test_pdf_merge_split_extract_preserves_page_order_and_content(tmp_path):
    a, b = tmp_path / "a.pdf", tmp_path / "b.pdf"
    _make_pdf(a, "ALPHA PAGE")
    _make_pdf(b, "BETA PAGE")
    merged = tmp_path / "merged.pdf"
    pdf.merge_pdfs([a, b], merged)
    assert len(_pdf_text(merged)) == 2
    assert "ALPHA PAGE" in _pdf_text(merged)[0]
    assert "BETA PAGE" in _pdf_text(merged)[1]
    parts = pdf.split_pdf_pages(merged, tmp_path)
    assert len(parts) == 2
    assert ["ALPHA PAGE" in _pdf_text(parts[0])[0], "BETA PAGE" in _pdf_text(parts[1])[0]] == [True, True]
    extracted = tmp_path / "extracted.pdf"
    pdf.extract_pdf_pages(merged, extracted, "2")
    assert len(_pdf_text(extracted)) == 1
    assert "BETA PAGE" in _pdf_text(extracted)[0]
    assert "ALPHA PAGE" not in _pdf_text(extracted)[0]


def test_pdf_delete_pages_keeps_selected_text(tmp_path):
    original = tmp_path / "original.pdf"
    doc = fitz.open()
    for label in ("ONE ONLY", "TWO ONLY", "THREE ONLY"):
        page = doc.new_page()
        page.insert_text((70, 70), label)
    doc.save(str(original))
    doc.close()
    result = tmp_path / "deleted.pdf"
    pdf.delete_pdf_pages(original, result, "2")
    texts = _pdf_text(result)
    assert len(texts) == 2
    assert "ONE ONLY" in texts[0] and "THREE ONLY" in texts[1]
    assert all("TWO ONLY" not in text for text in texts)


def test_word_to_pdf_preserves_editable_original_text_and_a4_page(tmp_path):
    original = tmp_path / "original.docx"
    document = Document()
    # The docx library defaults to US Letter. Select A4 explicitly and verify preservation.
    from docx.shared import Mm
    document.sections[0].page_width = Mm(210)
    document.sections[0].page_height = Mm(297)
    document.add_heading("INFINITY OFFICE QUALITY", level=1)
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip((table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)), ("Course", "Hours", "Physics", "03")):
        cell.text = value
    document.add_paragraph("Arabic line: اختبار الجودة")
    document.save(original)
    out = office.office_to_pdf(original, tmp_path / "rendered", timeout=75)
    pages = _pdf_text(out)
    assert len(pages) == 1, "Simple original DOCX unexpectedly overflowed to multiple PDF pages"
    joined = " ".join(pages)
    for expected in ("INFINITY OFFICE QUALITY", "Course", "Hours", "Physics", "03"):
        assert expected in joined
    with fitz.open(str(out)) as rendered:
        assert abs(rendered[0].rect.width - 595) < 5
        assert abs(rendered[0].rect.height - 842) < 5
        # Arabic glyph presence, not exact extraction order (PDF bidi varies).
        assert any("\u0600" <= char <= "\u06ff" for char in rendered[0].get_text("text"))


def test_pdf_to_word_roundtrip_retains_simple_page_content(tmp_path):
    original = tmp_path / "original.pdf"
    _make_pdf(original, "EDITABLE ROUNDTRIP", width=842, height=595)
    docx = tmp_path / "roundtrip.docx"
    mega_tools.pdf_to_docx(original, docx)
    editable = Document(str(docx))
    text = " ".join(p.text for p in editable.paragraphs)
    text += " " + " ".join(cell.text for table in editable.tables for row in table.rows for cell in row.cells)
    assert "EDITABLE ROUNDTRIP" in text
    assert abs(editable.sections[0].page_width.inches - 842 / 72) < 0.15
    converted = office.office_to_pdf(docx, tmp_path / "roundtrip-pdf", timeout=75)
    assert "EDITABLE ROUNDTRIP" in " ".join(_pdf_text(converted))
    with fitz.open(str(converted)) as rendered:
        assert len(rendered) <= 2
        assert rendered[0].rect.width > rendered[0].rect.height


def test_rgba_to_jpeg_background_and_geometry(tmp_path):
    original = tmp_path / "original.png"
    image = Image.new("RGBA", (90, 45), (255, 0, 0, 0))
    image.putpixel((45, 20), (0, 255, 0, 255))
    image.save(original)
    output = tmp_path / "out.jpg"
    images.convert_image(original, output, "JPEG")
    with Image.open(output) as converted:
        assert converted.format == "JPEG"
        assert converted.mode == "RGB"
        assert converted.size == (90, 45)
        assert sum(converted.getpixel((5, 5))) >= 680, "Transparent corner must not become black"


def test_resize_and_rotate_image_preserve_intended_dimensions(tmp_path):
    original = tmp_path / "wide.png"
    Image.new("RGB", (400, 200), (40, 120, 180)).save(original)
    resized = tmp_path / "resized.png"
    images.resize_image(original, resized, 100)
    with Image.open(resized) as result:
        assert result.size == (100, 50)
    rotated = tmp_path / "rotated.png"
    images.rotate_image(original, rotated, 90)
    with Image.open(rotated) as result:
        assert result.size == (200, 400)
    with pytest.raises(ValueError):
        images.rotate_image(original, tmp_path / "wrong.png", 13)


def test_invert_image_twice_is_lossless_for_rgba(tmp_path):
    original = tmp_path / "rgba.png"
    Image.new("RGBA", (37, 29), (12, 78, 201, 133)).save(original)
    once, twice = tmp_path / "once.png", tmp_path / "twice.png"
    mega_tools.image_invert(original, once)
    mega_tools.image_invert(once, twice)
    with Image.open(original) as a, Image.open(twice) as b:
        assert ImageChops.difference(a, b).getbbox() is None


def test_csv_json_csv_roundtrip_preserves_arabic_quotes_and_newlines(tmp_path):
    original = tmp_path / "in.csv"
    expected = [{"name": "اسم، مركب", "note": 'قال "نعم"\nالسطر الثاني', "code": "0012"},
                {"name": "English", "note": "comma, cell", "code": "0007"}]
    with original.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["name", "note", "code"])
        writer.writeheader()
        writer.writerows(expected)
    json_path = tmp_path / "out.json"
    office_advanced.csv_to_json(original, json_path)
    assert json.loads(json_path.read_text(encoding="utf-8")) == expected
    back = tmp_path / "back.csv"
    office_advanced.json_to_csv(json_path, back)
    with back.open("r", encoding="utf-8", newline="") as handle:
        assert list(csv.DictReader(handle)) == expected


def test_xlsx_to_json_preserves_multiple_sheets_and_types(tmp_path):
    original = tmp_path / "sheets.xlsx"
    workbook = Workbook()
    one = workbook.active
    one.title = "Arabic"
    one.append(["name", "amount", "active"])
    one.append(["أحمد", 123.25, True])
    two = workbook.create_sheet("Other")
    two.append(["id", "comment"])
    two.append(["0012", "مرحبا"])
    workbook.save(original)
    output = tmp_path / "sheets.json"
    office_advanced.xlsx_to_json(original, output)
    result = json.loads(output.read_text(encoding="utf-8"))
    assert list(result) == ["Arabic", "Other"]
    assert result["Arabic"] == [{"name": "أحمد", "amount": 123.25, "active": True}]
    assert result["Other"] == [{"id": "0012", "comment": "مرحبا"}]


def test_pptx_to_text_keeps_slide_order(tmp_path):
    original = tmp_path / "slides.pptx"
    pres = Presentation()
    for name in ("FIRST SLIDE", "SECOND SLIDE"):
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        box = slide.shapes.add_textbox(100000, 100000, 4000000, 500000)
        box.text = name
    pres.save(original)
    output = tmp_path / "slides.txt"
    office_advanced.pptx_to_text(original, output)
    result = output.read_text(encoding="utf-8")
    assert result.index("FIRST SLIDE") < result.index("SECOND SLIDE")
    assert "--- Slide 1 ---" in result and "--- Slide 2 ---" in result


@pytest.mark.parametrize("suffix,compress,decompress", [
    ("bz2", mega_tools.bzip2_compress, mega_tools.bzip2_decompress),
    ("xz", mega_tools.xz_compress, mega_tools.xz_decompress),
])
def test_archive_compress_decompress_is_bit_exact(tmp_path, suffix, compress, decompress):
    content = ("بيانات عربية\n" * 60).encode("utf-8") + bytes(range(256)) * 8
    original = tmp_path / "original.bin"
    original.write_bytes(content)
    packed, unpacked = tmp_path / ("archive." + suffix), tmp_path / "unpacked.bin"
    compress(original, packed)
    decompress(packed, unpacked)
    assert unpacked.read_bytes() == content


def test_urls_roundtrip_unicode_spaces_and_punctuation(tmp_path):
    original = tmp_path / "original.txt"
    message = "مرحباً / hello?x=1&name=أحمد"
    original.write_text(message, encoding="utf-8")
    encoded, decoded = tmp_path / "encoded.txt", tmp_path / "decoded.txt"
    mega_tools.url_encode(original, encoded)
    assert "%D9" in encoded.read_text(encoding="utf-8")
    mega_tools.url_decode(encoded, decoded)
    assert decoded.read_text(encoding="utf-8") == message


def test_pdf_to_text_retains_content_and_page_order(tmp_path):
    original = tmp_path / "two.pdf"
    doc = fitz.open()
    for label in ("PAGE A SOURCE", "PAGE B SOURCE"):
        page = doc.new_page()
        page.insert_text((60, 75), label)
    doc.save(str(original))
    doc.close()
    output = tmp_path / "text.txt"
    pdf.pdf_to_text(original, output)
    result = output.read_text(encoding="utf-8")
    assert result.index("PAGE A SOURCE") < result.index("PAGE B SOURCE")


def test_pdf_to_docx_rejects_empty_scanned_document(tmp_path):
    original = tmp_path / "blank.pdf"
    with fitz.open() as doc:
        doc.new_page()
        doc.save(str(original))
    output = tmp_path / "blank.docx"
    with pytest.raises(ValueError, match="OCR"):
        mega_tools.pdf_to_docx(original, output)
    assert not output.exists()
