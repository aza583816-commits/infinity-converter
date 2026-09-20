"""Independent PDF fidelity gates for larger Arabic, English, and bilingual inputs.

Unlike a one-file smoke test, these cases compare every distinctive source
record to *actual selectable output text*, verify page boundaries, and detect
blank/clipped pages. They intentionally fail on lossy renderers.
"""
from __future__ import annotations

import unicodedata
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

from converters.office import office_to_pdf


def normalized(text: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", text).split())


def pdf_text_and_geometry(path: Path) -> str:
    text = []
    with pymupdf.open(path) as pdf:
        assert len(pdf) > 0, "The converter produced an empty PDF"
        for page_no, page in enumerate(pdf, 1):
            content = page.get_text("text")
            assert content.strip(), ("Blank or unextractable PDF page", page_no)
            text.append(content)
            for word in page.get_text("words"):
                x0, y0, x1, y1 = word[:4]
                assert (-2 <= x0 < x1 <= page.rect.width + 2
                        and -2 <= y0 < y1 <= page.rect.height + 2), (
                    "PDF contains off-page/clipped selectable text", page_no, word[:5]
                )
    return normalized(" ".join(text))


@pytest.mark.parametrize("language", ("arabic", "english", "bilingual"))
def test_long_word_pdf_preserves_every_original_marker(tmp_path: Path, language: str):
    source = tmp_path / f"long-{language}.docx"
    doc = Document()
    doc.add_heading("INFINITY DOCUMENT TEST", 0)
    markers = []
    for page_index in range(1, 9):
        if page_index != 1:
            doc.add_page_break()
        for row_index in range(1, 7):
            marker = f"ID{page_index:02d}{row_index:02d}X"
            if language == "arabic":
                value = f"رقم {marker} سلامة الوقاية من الحريق"
            elif language == "english":
                value = f"Record {marker} fire protection safety"
            else:
                value = f"سجل {marker} fire protection السلامة"
            doc.add_paragraph(value)
            markers.append(value)
        table = doc.add_table(rows=3, cols=3)
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                tag = f"T{page_index:02d}{row_index}{col_index}X"
                value = (f"خلية {tag}" if language == "arabic"
                         else f"Cell {tag}" if language == "english"
                         else f"خلية Cell {tag}")
                cell.text = value
                markers.append(value)
    doc.save(source)
    output = office_to_pdf(source, tmp_path / "output", timeout=120)
    actual = pdf_text_and_geometry(output)
    missing = [value for value in markers if normalized(value) not in actual]
    assert not missing, ("Source Word content lost or reordered within fields", missing[:12])
    # Each distinct page's first paragraph must survive in original reading order.
    positions = [actual.find(normalized(markers[(n - 1) * 15])) for n in range(1, 9)]
    assert all(position >= 0 for position in positions)
    assert positions == sorted(positions), ("Word page content re-ordered", positions)


@pytest.mark.parametrize("language", ("arabic", "english", "bilingual"))
def test_large_two_sheet_excel_pdf_keeps_unique_values(tmp_path: Path, language: str):
    source = tmp_path / f"ledger-{language}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Primary"
    ws.append(["id", "name", "amount"])
    expected = []
    for i in range(1, 201):
        unique = f"Q{i:05d}Z"
        name = (f"السلامة {unique}" if language == "arabic" else
                f"Safety {unique}" if language == "english" else
                f"السلامة Safety {unique}")
        ws.append([unique, name, i + .125])
        expected.extend([unique, name])
    second = wb.create_sheet("Secondary")
    second.append(["id", "note"])
    for i in range(1, 51):
        marker = f"SECOND{i:04d}Z"
        second.append([marker, "السلامة" if language != "english" else "Safety"])
        expected.append(marker)
    wb.save(source)
    output = office_to_pdf(source, tmp_path / "output", timeout=120)
    actual = pdf_text_and_geometry(output)
    missing = [value for value in expected if normalized(value) not in actual]
    assert not missing, ("Excel-to-PDF dropped source cells", missing[:12])
