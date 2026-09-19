"""Interleaved paragraph/table content must retain original reading order."""
from docx import Document
from converters.office_advanced import docx_to_text, docx_to_html


def test_docx_text_and_html_preserve_paragraph_table_order(tmp_path):
    source = tmp_path / "mixed.docx"
    doc = Document()
    doc.add_paragraph("Before table")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "اسم"
    doc.add_paragraph("After table")
    doc.save(source)
    text_file, html_file = tmp_path / "mixed.txt", tmp_path / "mixed.html"
    docx_to_text(source, text_file)
    docx_to_html(source, html_file)
    txt = text_file.read_text(encoding="utf-8")
    html = html_file.read_text(encoding="utf-8")
    assert txt.index("Before table") < txt.index("Name\tاسم") < txt.index("After table")
    assert html.index("Before table") < html.index("<table>") < html.index("After table")
    assert "Name" in html and "اسم" in html
