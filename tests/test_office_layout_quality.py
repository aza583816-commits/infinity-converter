"""Quality regression for the actual Word-to-PDF renderer and PDF-to-DOCX roundtrip.

Uses only a generated synthetic document. Output quality is measured via page
geometry, text retention, and non-clipped word positions; these checks do not
claim pixel-perfect PDF-to-editable-Word fidelity for all PDFs.
"""
from __future__ import annotations

import shutil

import pymupdf
import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt

from converters.mega_tools import pdf_to_docx
from converters.office import office_to_pdf


pytestmark = pytest.mark.skipif(
    shutil.which("libreoffice") is None,
    reason="Real LibreOffice rendering is required for the layout-quality gate",
)


def _fixture(path):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    doc.add_heading("BILINGUAL COURSE SCHEDULE", 1)
    doc.add_paragraph("السلامة والوقاية من الحريق — Safety and Fire Protection")
    table = doc.add_table(rows=4, cols=3)
    for row, cells in enumerate((
        ("Course", "Time", "Room"),
        ("Electrical Circuits", "09:00", "E074"),
        ("Physics Lab", "11:00", "603"),
        ("Chemistry", "13:20", "511"),
    )):
        for index, value in enumerate(cells):
            table.cell(row, index).text = value
    doc.add_paragraph("FINAL CHECK: All course rows must remain readable.")
    doc.save(path)


def _text(docx):
    word = Document(str(docx))
    parts = [p.text for p in word.paragraphs]
    for table in word.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def test_word_to_pdf_retains_page_geometry_content_and_visible_table(tmp_path):
    original = tmp_path / "course-schedule.docx"
    output_dir = tmp_path / "rendered"
    _fixture(original)
    converted = office_to_pdf(original, output_dir, timeout=90)
    with pymupdf.open(str(converted)) as pdf:
        assert len(pdf) == 1, "Compact landscape table unexpectedly overflowed"
        page = pdf[0]
        assert page.rect.width > page.rect.height, "Landscape orientation was lost"
        assert abs(page.rect.width - 11.69 * 72) < 4
        assert abs(page.rect.height - 8.27 * 72) < 4
        content = page.get_text()
        for token in ("BILINGUAL COURSE SCHEDULE", "Electrical Circuits", "Physics Lab", "Chemistry", "E074", "603", "511", "FINAL CHECK"):
            assert token in content, f"Word-to-PDF lost {token!r}"
        for word in page.get_text("words"):
            x0, y0, x1, y1 = word[:4]
            assert -1 <= x0 < x1 <= page.rect.width + 1
            assert -1 <= y0 < y1 <= page.rect.height + 1


def test_pdf_to_word_roundtrip_retains_editable_information_and_landscape(tmp_path):
    original = tmp_path / "course-schedule.docx"
    _fixture(original)
    source_pdf = office_to_pdf(original, tmp_path / "first-pdf", timeout=90)
    reconstructed = tmp_path / "reconstructed.docx"
    pdf_to_docx(source_pdf, reconstructed)
    editable = _text(reconstructed)
    for token in ("BILINGUAL COURSE SCHEDULE", "Electrical Circuits", "Physics Lab", "Chemistry", "09:00", "E074", "603", "FINAL CHECK"):
        assert token in editable, f"PDF-to-DOCX omitted {token!r}"
    result_word = Document(str(reconstructed))
    assert result_word.sections
    assert result_word.sections[0].page_width > result_word.sections[0].page_height
    final_pdf = office_to_pdf(reconstructed, tmp_path / "roundtrip-pdf", timeout=90)
    with pymupdf.open(str(final_pdf)) as pdf:
        assert 1 <= len(pdf) <= 2, "Roundtrip expanded a one-page table excessively"
        all_text = "\n".join(page.get_text() for page in pdf)
        for token in ("BILINGUAL COURSE SCHEDULE", "Electrical Circuits", "Physics Lab", "Chemistry", "FINAL CHECK"):
            assert token in all_text, f"Roundtrip PDF omitted {token!r}"


def _assert_table_columns_and_rows(page):
    """Validate the rendered table's visual geometry, not only text presence."""
    expected_rows = [
        ("Course", "Time", "Room"),
        ("Electrical Circuits", "09:00", "E074"),
        ("Physics Lab", "11:00", "603"),
        ("Chemistry", "13:20", "511"),
    ]
    for labels in expected_rows:
        bounds = []
        for label in labels:
            regions = page.search_for(label)
            assert regions, f"Table cell {label!r} is missing in rendered PDF"
            rect = regions[0]
            assert rect.x0 >= -1 and rect.y0 >= -1, f"Table cell {label!r} starts off-page"
            assert rect.x1 <= page.rect.width + 1 and rect.y1 <= page.rect.height + 1, (
                f"Table cell {label!r} is clipped by page boundary"
            )
            bounds.append(rect)
        assert bounds[0].x0 < bounds[1].x0 < bounds[2].x0, (
            f"Table columns are no longer ordered on the page: {labels}"
        )
        assert max(rect.y0 for rect in bounds) - min(rect.y0 for rect in bounds) < 24, (
            f"Table cells have slipped into separate visual rows: {labels}"
        )


def test_editable_pdf_to_word_roundtrip_retains_visible_table_grid(tmp_path):
    original = tmp_path / "source.docx"
    _fixture(original)
    source_pdf = office_to_pdf(original, tmp_path / "source-pdf", timeout=90)
    reconstructed = tmp_path / "reconstructed.docx"
    pdf_to_docx(source_pdf, reconstructed)
    result = office_to_pdf(reconstructed, tmp_path / "final-pdf", timeout=90)
    with pymupdf.open(str(result)) as doc:
        assert len(doc) == 1, "One-page table spilled over multiple pages on roundtrip"
        _assert_table_columns_and_rows(doc[0])
