import io
import zipfile
from pathlib import Path

import pytest
from pypdf import PdfWriter

from core.storage import TempWorkspace
from converters.office import _reject_remote_html_resources
from security.file_guard import validate_upload


class Upload:
    def __init__(self, name: str, data: bytes, content_type: str = "application/octet-stream"):
        self.filename = name
        self.mimetype = content_type
        self._stream = io.BytesIO(data)

    def read(self, n=-1):
        return self._stream.read(n)


def fs(name: str, data: bytes, content_type: str = "application/octet-stream") -> Upload:
    return Upload(name, data, content_type)


def test_temp_workspace_containment_and_idempotent_cleanup():
    workspace = TempWorkspace().__enter__()
    root = workspace.path
    assert root and root.exists()
    assert workspace.contains_input(workspace.input_dir / "a.txt")
    assert workspace.contains_output(workspace.output_dir / "a.txt")
    assert not workspace.contains_input(Path("/tmp/outside.txt"))
    workspace.cleanup()
    workspace.cleanup()
    assert not root.exists()


def test_upload_filename_is_reduced_to_basename_and_saved_under_input():
    with TempWorkspace() as workspace:
        result = validate_upload(
            fs("../../notes.txt", b"hello\n", "text/plain"),
            max_bytes=1024,
            inspect_only=False,
            workspace=workspace.path,
        )
        assert result["filename"] == "notes.txt"
        assert result["safe"] is True
        assert workspace.contains_input(result["path"])


def test_signature_mismatch_is_rejected():
    with pytest.raises(ValueError, match="توقيع"):
        validate_upload(
            fs("fake.pdf", b"not a pdf", "application/pdf"),
            max_bytes=1024,
            inspect_only=True,
        )


def test_zip_symlink_is_rejected():
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as zf:
        info = zipfile.ZipInfo("link")
        info.create_system = 3
        info.external_attr = (0o120777 << 16)
        zf.writestr(info, "target")
    with pytest.raises(ValueError, match="رابط"):
        validate_upload(
            fs("links.zip", payload.getvalue(), "application/zip"),
            max_bytes=1024 * 1024,
            inspect_only=True,
        )


def test_encrypted_pdf_can_reach_unlock_tool_validation_stage():
    payload = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("secret")
    writer.write(payload)
    result = validate_upload(
        fs("locked.pdf", payload.getvalue(), "application/pdf"),
        max_bytes=1024 * 1024,
        inspect_only=True,
    )
    assert result["safe"] is True
    assert result["encrypted"] is True


def test_ooxml_external_embedded_image_is_rejected():
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", "<document/>")
        zf.writestr(
            "word/_rels/document.xml.rels",
            '<Relationships><Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" TargetMode="External" Target="https://example.com/a.png"/></Relationships>',
        )
    with pytest.raises(ValueError, match="مورد خارجي"):
        validate_upload(
            fs(
                "external.docx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
            max_bytes=1024 * 1024,
            inspect_only=True,
        )


def test_html_conversion_rejects_remote_render_resources(tmp_path):
    source = tmp_path / "remote.html"
    source.write_text('<html><body><img src="https://example.com/a.png"></body></html>', encoding="utf-8")
    with pytest.raises(ValueError, match="الخصوصية"):
        _reject_remote_html_resources(source)


def test_spreadsheet_exports_neutralize_formula_injection(tmp_path):
    from openpyxl import load_workbook
    from converters.office import csv_to_xlsx
    from converters.office_advanced import json_to_xlsx

    csv_source = tmp_path / "danger.csv"
    csv_source.write_text("name,value\nAlice,=1+1\nBob,@SUM(A1:A2)\n", encoding="utf-8")
    csv_output = tmp_path / "safe-csv.xlsx"
    csv_to_xlsx(csv_source, csv_output)
    workbook = load_workbook(csv_output, data_only=False)
    sheet = workbook.active
    assert sheet["B2"].value == "'=1+1"
    assert sheet["B3"].value == "'@SUM(A1:A2)"
    workbook.close()

    json_source = tmp_path / "danger.json"
    json_source.write_text('[{"value":"+1+1"},{"value":"-2+3"}]', encoding="utf-8")
    json_output = tmp_path / "safe-json.xlsx"
    json_to_xlsx(json_source, json_output)
    workbook = load_workbook(json_output, data_only=False)
    sheet = workbook.active
    assert sheet["A2"].value == "'+1+1"
    assert sheet["A3"].value == "'-2+3"
    workbook.close()
