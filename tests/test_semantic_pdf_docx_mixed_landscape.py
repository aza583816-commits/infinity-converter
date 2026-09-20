"""Mixed-orientation PDF must preserve its fourth page and unique source values in DOCX."""
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.shared import Pt

from scripts.semantic_output_audit import independent_oracle


def _source(path: Path) -> None:
    with pymupdf.open() as pdf:
        for index in range(1, 4):
            page = pdf.new_page(width=595, height=842)
            page.insert_text((40, 60), f"INFINITY CONVERTER Page {index}", fontsize=12)
            page.insert_text((40, 90), "Email test@example.com Total 1250.50", fontsize=12)
        landscape = pdf.new_page(width=842, height=595)
        landscape.insert_text((40, 60), "LANDSCAPE PAGE FOUR", fontsize=14)
        landscape.insert_text((40, 90), "Columns Alpha 00123 | Beta 00999", fontsize=12)
        pdf.save(path)


def _output(path: Path, *, include_landscape: bool, include_fields: bool = True) -> None:
    doc = Document()
    doc.sections[0].page_width = Pt(595)
    for index in range(1, 4):
        doc.add_paragraph(f"INFINITY CONVERTER Page {index}")
    doc.add_paragraph("Email test@example.com Total 1250.50")
    if include_landscape:
        doc.add_paragraph("LANDSCAPE PAGE FOUR")
        if include_fields:
            doc.add_paragraph("Columns Alpha 00123 | Beta 00999")
    doc.save(path)


def test_mixed_pdf_fourth_landscape_page_and_fields_are_preserved(tmp_path: Path) -> None:
    source, result = tmp_path / "mixed.pdf", tmp_path / "converted.docx"
    _source(source)
    _output(result, include_landscape=True)
    assert independent_oracle("pdf-to-docx", source, result, tmp_path)


def test_missing_landscape_page_fails_independent_verification(tmp_path: Path) -> None:
    source, result = tmp_path / "mixed.pdf", tmp_path / "converted.docx"
    _source(source)
    _output(result, include_landscape=False)
    with pytest.raises(AssertionError, match="source page missing"):
        independent_oracle("pdf-to-docx", source, result, tmp_path)


def test_missing_landscape_source_cell_fails_independent_verification(tmp_path: Path) -> None:
    source, result = tmp_path / "mixed.pdf", tmp_path / "converted.docx"
    _source(source)
    _output(result, include_landscape=True, include_fields=False)
    with pytest.raises(AssertionError, match="source field lost"):
        independent_oracle("pdf-to-docx", source, result, tmp_path)
