"""Real DOCX-to-PDF output must preserve complete Arabic/numeric table fields."""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

import pymupdf
from docx import Document
from pypdf import PdfReader

from converters.office import office_to_pdf


def _script_boundaries(value: str) -> str:
    value = " ".join(unicodedata.normalize("NFKC", value).split())
    return re.sub(
        r"(?<=[\u0600-\u06FF])(?=[0-9])|(?<=[0-9])(?=[\u0600-\u06FF])",
        " ", value,
    )


def test_real_arabic_word_table_has_complete_ids_and_visible_word_spacing(tmp_path: Path):
    source = tmp_path / "safety-arabic.docx"
    document = Document()
    document.add_paragraph("سجل السلامة Fire protection")
    table = document.add_table(rows=1, cols=3)
    ids = ("00123", "00124", "00999")
    values = [f"cell-{index} / قيمة {code}" for index, code in enumerate(ids)]
    for cell, value in zip(table.rows[0].cells, values):
        cell.text = value
    document.add_paragraph("AFTER TABLE END MARKER")
    document.save(source)

    output = office_to_pdf(source, tmp_path / "pdf", timeout=45)
    logical = " ".join(page.extract_text() or "" for page in PdfReader(str(output)).pages)
    actual = _script_boundaries(logical)
    assert "سجل السلامة Fire protection" in actual
    assert "AFTER TABLE END MARKER" in actual
    for value in values:
        expected = _script_boundaries(value)
        assert re.search(r"(?<!\w)" + re.escape(expected) + r"(?!\w)", actual), (
            "Missing or truncated Word table cell", value, actual[:700]
        )
    # A prefix of a different field must not be accepted as its original value.
    assert not re.search(r"(?<!\w)" + re.escape("cell-0 / قيمة 0012") + r"(?!\w)", actual)

    with pymupdf.open(output) as document:
        assert len(document) >= 1
        words = [word for page in document for word in page.get_text("words")]
        for code in ids:
            matching = [w for w in words if w[4] == code]
            assert len(matching) == 1, ("Lost or duplicated complete identifier", code)
            number = matching[0]
            arabic = [
                w for w in words
                if w[4] == "قيمة"
                and abs(w[1] - number[1]) <= 2
                and 0 < min(abs(w[0] - number[2]), abs(number[0] - w[2])) < 20
            ]
            assert arabic, ("Arabic word and ID overlap or have no visible gap", code)
