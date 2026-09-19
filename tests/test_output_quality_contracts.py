"""Content-level conversion assertions using synthetic, non-sensitive fixtures.

These tests are intentionally more demanding than the 162-operation smoke tests:
they assert decoded values, coordinates, page geometry, color and data retention.
They do not claim universal fidelity for arbitrary input files.
"""
import csv
import io
import json
import zipfile

import fitz
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image

from converters import office_advanced, images, archive
from converters.office import office_to_pdf


def test_csv_json_csv_preserves_quoted_unicode_and_row_order(tmp_path):
    original = tmp_path / "input.csv"
    records = [
        {"name": "مرحبا, World", "note": 'quote "value"', "value": "0007"},
        {"name": "Second", "note": "line one\\nline two", "value": "3.14"},
    ]
    with original.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    json_file = tmp_path / "output.json"
    csv_file = tmp_path / "roundtrip.csv"
    office_advanced.csv_to_json(original, json_file)
    assert json.loads(json_file.read_text(encoding="utf-8")) == records
    office_advanced.json_to_csv(json_file, csv_file)
    with csv_file.open(encoding="utf-8", newline="") as stream:
        assert list(csv.DictReader(stream)) == records


def test_xlsx_json_preserves_multisheet_rows_unicode_and_numbers(tmp_path):
    workbook = Workbook()
    first = workbook.active
    first.title = "جدول"
    first.append(["name", "score", "optional"])
    first.append(["طالب", 42, None])
    second = workbook.create_sheet("Second")
    second.append(["item", "price"])
    second.append(["Book", 12.5])
    source = tmp_path / "sheets.xlsx"
    workbook.save(source)
    result = tmp_path / "sheets.json"
    office_advanced.xlsx_to_json(source, result)
    decoded = json.loads(result.read_text(encoding="utf-8"))
    assert decoded["جدول"] == [{"name": "طالب", "score": 42, "optional": None}]
    assert decoded["Second"] == [{"item": "Book", "price": 12.5}]


def test_docx_html_escapes_user_text_and_preserves_table_values(tmp_path):
    source = tmp_path / "source.docx"
    word = Document()
    word.add_paragraph("Hello <script>alert(1)</script>")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "عربي"
    table.cell(0, 1).text = "A & B"
    word.save(source)
    result = tmp_path / "result.html"
    office_advanced.docx_to_html(source, result)
    html = result.read_text(encoding="utf-8")
    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    assert "<table>" in html and "عربي" in html and "A &amp; B" in html


def test_image_to_jpg_preserves_dimensions_and_composites_transparency(tmp_path):
    image = Image.new("RGBA", (48, 32), (255, 0, 0, 0))
    image.putpixel((12, 12), (0, 0, 255, 255))
    source = tmp_path / "alpha.png"
    image.save(source)
    result = tmp_path / "alpha.jpg"
    images.convert_image(source, result, "JPEG")
    with Image.open(result) as converted:
        assert converted.format == "JPEG"
        assert converted.size == (48, 32)
        assert converted.mode == "RGB"
        white = converted.getpixel((40, 25))
        blue = converted.getpixel((12, 12))
        assert min(white) > 238
        assert blue[2] > blue[0] and blue[2] > blue[1]


def test_image_rotation_has_expected_geometry_and_pixel_regions(tmp_path):
    source = tmp_path / "marker.png"
    image = Image.new("RGB", (30, 20), (255, 255, 255))
    for x in range(0, 10):
        for y in range(0, 10):
            image.putpixel((x, y), (0, 0, 0))
    image.save(source)
    result = tmp_path / "rotated.png"
    images.rotate_image(source, result, 90)
    with Image.open(result) as rotated:
        assert rotated.size == (20, 30)
        assert rotated.getpixel((15, 5)) == (0, 0, 0)
        assert rotated.getpixel((5, 25)) == (255, 255, 255)


def test_zip_create_extract_preserves_unicode_names_and_exact_bytes(tmp_path):
    source = tmp_path / "عربي.txt"
    payload = "Unicode content اختبار\\n".encode("utf-8")
    source.write_bytes(payload)
    output = tmp_path / "archive.zip"
    archive.create_zip([(source, "folder/عربي.txt")], output)
    with zipfile.ZipFile(output) as zf:
        assert zf.read("folder/عربي.txt") == payload
    extracted = archive.extract_zip(output, tmp_path / "unpacked")
    assert len(extracted) == 1
    assert extracted[0].name == "عربي.txt"
    assert extracted[0].read_bytes() == payload


@pytest.mark.skipif(__import__("shutil").which("libreoffice") is None, reason="Office engine missing")
def test_word_to_pdf_preserves_landscape_geometry_text_and_table(tmp_path):
    source = tmp_path / "landscape.docx"
    word = Document()
    section = word.sections[0]
    section.page_width, section.page_height = section.page_height, section.page_width
    word.add_heading("DOCUMENT QUALITY CHECK", 1)
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Course"
    table.cell(0, 1).text = "Grade"
    table.cell(1, 0).text = "Physics"
    table.cell(1, 1).text = "95"
    word.save(source)
    rendered = office_to_pdf(source, tmp_path / "output", timeout=90)
    with fitz.open(rendered) as pdf:
        assert len(pdf) == 1
        assert pdf[0].rect.width > pdf[0].rect.height
        txt = pdf[0].get_text("text")
        for marker in ("DOCUMENT QUALITY CHECK", "Course", "Grade", "Physics", "95"):
            assert marker in txt, (marker, txt[:600])
        blocks = pdf[0].get_text("blocks")
        assert all(
            rect[0] >= -2 and rect[1] >= -2 and rect[2] <= pdf[0].rect.width + 2
            and rect[3] <= pdf[0].rect.height + 2
            for rect in (b[:4] for b in blocks)
        )
